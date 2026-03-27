from __future__ import annotations

from io import BytesIO

from openpyxl import Workbook
from openpyxl.chart import ScatterChart, Series, Reference

from rc_bending.materials import MaterialCatalog
from rc_bending.models import BendingResult, CurvePoint, SectionInput, ServiceabilityReport
from rc_bending.serviceability import calculate_deflection_mm_from_curvature
from rc_bending.solver import build_layer_force_table, build_strain_profile_for_point


def _add_moment_curvature_chart(sheet) -> None:
    chart = ScatterChart()
    chart.title = "Moment-Curvature"
    chart.x_axis.title = "Curvature [1/m]"
    chart.y_axis.title = "Moment [kN m]"
    x_values = Reference(sheet, min_col=5, min_row=2, max_row=sheet.max_row)
    y_values = Reference(sheet, min_col=4, min_row=1, max_row=sheet.max_row)
    series = Series(y_values, x_values, title_from_data=True)
    chart.series.append(series)
    chart.height = 8
    chart.width = 14
    sheet.add_chart(chart, "H2")


def _add_strain_profile_chart(sheet) -> None:
    chart = ScatterChart()
    chart.title = "Concrete Strain Profile"
    chart.x_axis.title = "Strain [-]"
    chart.y_axis.title = "Depth z [mm]"
    x_values = Reference(sheet, min_col=2, min_row=2, max_row=sheet.max_row)
    y_values = Reference(sheet, min_col=1, min_row=1, max_row=sheet.max_row)
    series = Series(y_values, x_values, title_from_data=True)
    chart.series.append(series)
    chart.height = 8
    chart.width = 10
    sheet.add_chart(chart, "E2")


def _add_xy_scatter_chart(
    sheet,
    *,
    title: str,
    x_axis_title: str,
    y_axis_title: str,
    x_col: int,
    y_col: int,
    anchor: str,
    width: float = 14,
    height: float = 8,
) -> None:
    chart = ScatterChart()
    chart.title = title
    chart.x_axis.title = x_axis_title
    chart.y_axis.title = y_axis_title
    x_values = Reference(sheet, min_col=x_col, min_row=2, max_row=sheet.max_row)
    y_values = Reference(sheet, min_col=y_col, min_row=1, max_row=sheet.max_row)
    series = Series(y_values, x_values, title_from_data=True)
    chart.series.append(series)
    chart.height = height
    chart.width = width
    sheet.add_chart(chart, anchor)


def _to_strain_e5(value: float) -> float:
    return value * 100000.0


def _strain_at_depth(
    *,
    top_strain: float,
    bottom_strain: float,
    depth_mm: float,
    section_height_mm: float,
) -> float:
    if abs(section_height_mm) <= 1e-9:
        return top_strain
    return top_strain + (bottom_strain - top_strain) * (depth_mm / section_height_mm)


def _rebar_display_sign(section: SectionInput, *, rebar_index: int) -> float:
    indexed_layers = list(enumerate(section.rebar_layers, start=1))
    if len(indexed_layers) == 1:
        return -1.0
    bottom_rebar_index = max(indexed_layers, key=lambda item: item[1].z_mm)[0]
    return -1.0 if rebar_index == bottom_rebar_index else 1.0


def _build_rebar_strain_values_e5(
    section: SectionInput,
    result: BendingResult,
    *,
    rebar_index: int,
) -> list[float]:
    rebar = section.rebar_layers[rebar_index - 1]
    display_sign = _rebar_display_sign(section, rebar_index=rebar_index)
    return [
        display_sign
        * _to_strain_e5(
            _strain_at_depth(
                top_strain=point.top_strain,
                bottom_strain=point.bottom_strain,
                depth_mm=rebar.z_mm,
                section_height_mm=section.section_height_mm,
            )
        )
        for point in result.curve_points
    ]


def _append_sheet_rows(sheet, rows: list[list[float | int | str]]) -> None:
    for row in rows:
        sheet.append(row)


def _add_concrete_theory_sheet(workbook: Workbook, result: BendingResult) -> None:
    sheet = workbook.create_sheet("ConcreteMomentStrainTheory")
    rows = sorted(
        [
            [point.step_index, point.moment_kNm, _to_strain_e5(point.top_strain)]
            for point in result.curve_points
        ],
        key=lambda row: (float(row[2]), int(row[0])),
    )
    sheet.append(["Крок", "M, кН·м", "ε_c,top, 10^-5"])
    _append_sheet_rows(sheet, rows)
    _add_xy_scatter_chart(
        sheet,
        title="Concrete Moment-Strain Theory",
        x_axis_title="ε_c,top [10^-5]",
        y_axis_title="Moment [kN m]",
        x_col=3,
        y_col=2,
        anchor="E2",
    )


def _add_rebar_theory_sheets(workbook: Workbook, section: SectionInput, result: BendingResult) -> None:
    indexed_layers = list(enumerate(section.rebar_layers, start=1))
    if len(indexed_layers) == 1:
        rebar_index = indexed_layers[0][0]
        rows = sorted(
            [
                [point.step_index, point.moment_kNm, strain_e5]
                for point, strain_e5 in zip(
                    result.curve_points,
                    _build_rebar_strain_values_e5(section, result, rebar_index=rebar_index),
                    strict=True,
                )
            ],
            key=lambda row: (float(row[2]), int(row[0])),
        )
        sheet = workbook.create_sheet("RebarMomentStrainTheory")
        sheet.append(["Крок", "M, кН·м", "ε_s, 10^-5"])
        _append_sheet_rows(sheet, rows)
        _add_xy_scatter_chart(
            sheet,
            title="Rebar Moment-Strain Theory",
            x_axis_title="ε_s [10^-5]",
            y_axis_title="Moment [kN m]",
            x_col=3,
            y_col=2,
            anchor="E2",
        )
        return

    top_rebar_index = min(indexed_layers, key=lambda item: item[1].z_mm)[0]
    bottom_rebar_index = max(indexed_layers, key=lambda item: item[1].z_mm)[0]
    sheet_configs = [
        ("TopRebarMomentStrainTheory", top_rebar_index, "ε_s,top, 10^-5", "Top Rebar Moment-Strain Theory"),
        ("BottomRebarMomentStrainTheory", bottom_rebar_index, "ε_s,bot, 10^-5", "Bottom Rebar Moment-Strain Theory"),
    ]
    for sheet_name, rebar_index, strain_label, chart_title in sheet_configs:
        rows = sorted(
            [
                [point.step_index, point.moment_kNm, strain_e5]
                for point, strain_e5 in zip(
                    result.curve_points,
                    _build_rebar_strain_values_e5(section, result, rebar_index=rebar_index),
                    strict=True,
                )
            ],
            key=lambda row: (float(row[2]), int(row[0])),
        )
        sheet = workbook.create_sheet(sheet_name)
        sheet.append(["Крок", "M, кН·м", strain_label])
        _append_sheet_rows(sheet, rows)
        _add_xy_scatter_chart(
            sheet,
            title=chart_title,
            x_axis_title=f"{strain_label} [10^-5]",
            y_axis_title="Moment [kN m]",
            x_col=3,
            y_col=2,
            anchor="E2",
        )


def _add_deflection_theory_sheet(workbook: Workbook, result: BendingResult, serviceability_report: ServiceabilityReport) -> None:
    service_input = serviceability_report.input
    rows = [
        [
            point.step_index,
            point.moment_kNm,
            calculate_deflection_mm_from_curvature(
                curvature_1_per_m=point.curvature_1_per_m,
                span_mm=service_input.span_mm,
                support_scheme=service_input.support_scheme,
                a_mm=service_input.a_mm,
                phi_creep=service_input.phi_creep,
            ),
        ]
        for point in result.curve_points
    ]
    sheet = workbook.create_sheet("DeflectionCurveTheory")
    sheet.append(["Крок", "M, кН·м", "f, мм"])
    _append_sheet_rows(sheet, rows)
    _add_xy_scatter_chart(
        sheet,
        title="Deflection Theory",
        x_axis_title="f [mm]",
        y_axis_title="Moment [kN m]",
        x_col=3,
        y_col=2,
        anchor="E2",
    )


def build_results_workbook(
    section: SectionInput,
    catalog: MaterialCatalog,
    result: BendingResult,
    *,
    selected_point: CurvePoint | None = None,
    serviceability_report: ServiceabilityReport | None = None,
) -> Workbook:
    export_point = selected_point or result.peak_point
    workbook = Workbook()
    inputs_sheet = workbook.active
    inputs_sheet.title = "Inputs"
    inputs_sheet.append(["Parameter", "Value"])
    inputs_sheet.append(["section_height_mm", section.section_height_mm])
    for index, layer in enumerate(section.concrete_layers, start=1):
        inputs_sheet.append([f"concrete_{index}_class", layer.concrete_class])
        inputs_sheet.append([f"concrete_{index}_width_mm", layer.width_mm])
        inputs_sheet.append([f"concrete_{index}_height_mm", layer.height_mm])
    for index, layer in enumerate(section.rebar_layers, start=1):
        inputs_sheet.append([f"rebar_{index}_z_mm", layer.z_mm])
        inputs_sheet.append([f"rebar_{index}_area_mm2", layer.area_mm2])
        inputs_sheet.append([f"rebar_{index}_steel_class", layer.steel_class])

    materials_sheet = workbook.create_sheet("Materials")
    materials_sheet.append(["Type", "Class", "f_cd/f_yk", "E", "epsilon_limit", "extra"])
    seen_concrete = set()
    for layer in section.concrete_layers:
        if layer.concrete_class in seen_concrete:
            continue
        material = catalog.concrete[layer.concrete_class]
        materials_sheet.append(
            ["concrete", material.concrete_class, material.f_cd_mpa, material.e_cd_gpa, material.epsilon_cu1, ",".join(str(x) for x in material.a)]
        )
        seen_concrete.add(layer.concrete_class)
    seen_steel = set()
    for layer in section.rebar_layers:
        if layer.steel_class in seen_steel:
            continue
        material = catalog.steel[layer.steel_class]
        materials_sheet.append(
            ["steel", material.steel_class, material.f_yk_mpa, material.e_s_mpa, material.epsilon_ud, material.gamma_s]
        )
        seen_steel.add(layer.steel_class)

    curve_sheet = workbook.create_sheet("MomentCurvature")
    curve_sheet.append(
        ["step_index", "top_strain", "bottom_strain", "moment_kNm", "curvature_1_per_m", "axial_residual_kN", "neutral_axis_mm", "state_label"]
    )
    for point in result.curve_points:
        curve_sheet.append(
            [
                point.step_index,
                point.top_strain,
                point.bottom_strain,
                point.moment_kNm,
                point.curvature_1_per_m,
                point.axial_residual_kN,
                point.neutral_axis_mm,
                point.state_label,
            ]
        )
    _add_moment_curvature_chart(curve_sheet)

    strain_sheet = workbook.create_sheet("ConcreteStrainProfile")
    strain_sheet.append(["z_mm", "strain"])
    for profile_point in build_strain_profile_for_point(section, export_point):
        strain_sheet.append([profile_point.z_mm, profile_point.strain])
    _add_strain_profile_chart(strain_sheet)

    iteration_sheet = workbook.create_sheet("IntermediateIterations")
    iteration_sheet.append(
        ["outer_step", "iteration", "lower_bottom_strain", "upper_bottom_strain", "trial_bottom_strain", "axial_residual_kN"]
    )
    for row in result.inner_iterations:
        iteration_sheet.append(
            [
                row.outer_step,
                row.iteration,
                row.lower_bottom_strain,
                row.upper_bottom_strain,
                row.trial_bottom_strain,
                row.axial_residual_kN,
            ]
        )

    layer_sheet = workbook.create_sheet("LayerForces")
    layer_sheet.append(["kind", "index", "class", "z_mm", "area_mm2", "strain", "stress_mpa", "force_kN"])
    for row in build_layer_force_table(section, catalog, export_point):
        layer_sheet.append(
            [
                row["kind"],
                row["index"],
                row["class"],
                row["z_mm"],
                row["area_mm2"],
                row["strain"],
                row["stress_mpa"],
                row["force_kN"],
            ]
        )

    _add_concrete_theory_sheet(workbook, result)
    _add_rebar_theory_sheets(workbook, section, result)

    if serviceability_report is not None:
        service_input = serviceability_report.input
        snapshot = serviceability_report.snapshot
        crack_width = serviceability_report.crack_width
        deflection = serviceability_report.deflection

        service_input_sheet = workbook.create_sheet("ServiceabilityInputs")
        service_input_sheet.append(["Parameter", "Value"])
        service_input_sheet.append(["span_mm", service_input.span_mm])
        service_input_sheet.append(["support_scheme", service_input.support_scheme])
        service_input_sheet.append(["a_mm", service_input.a_mm])
        service_input_sheet.append(["phi_creep", service_input.phi_creep])
        service_input_sheet.append(["deflection_limit_profile", service_input.deflection_limit_profile])
        service_input_sheet.append(["available_gap_mm", service_input.available_gap_mm])
        service_input_sheet.append(["w_limit_mm", service_input.w_limit_mm])
        service_input_sheet.append(["load_duration", service_input.load_duration])
        service_input_sheet.append(["selected_step_index", snapshot.step_index])
        service_input_sheet.append(["selected_moment_kNm", snapshot.moment_kNm])
        service_input_sheet.append(["selected_curvature_1_per_m", snapshot.curvature_1_per_m])

        crack_sheet = workbook.create_sheet("CrackWidthCheck")
        crack_sheet.append(["Parameter", "Value"])
        crack_sheet.append(["sigma_s_mpa", crack_width.sigma_s_mpa])
        crack_sheet.append(["rho_p_eff", crack_width.rho_p_eff])
        crack_sheet.append(["crack_spacing_mm", crack_width.crack_spacing_mm])
        crack_sheet.append(["strain_difference", crack_width.strain_difference])
        crack_sheet.append(["w_k_mm", crack_width.w_k_mm])
        crack_sheet.append(["w_limit_mm", crack_width.w_limit_mm])
        crack_sheet.append(["is_within_limit", "OK" if crack_width.is_within_limit else "NG"])
        crack_sheet.append(["derived_bar_spacing_mm", crack_width.derived_bar_spacing_mm])
        crack_sheet.append(["note", crack_width.note])

        deflection_sheet = workbook.create_sheet("DeflectionCheck")
        deflection_sheet.append(["Parameter", "Value"])
        deflection_sheet.append(["curvature_1_per_m", deflection.curvature_1_per_m])
        deflection_sheet.append(["effective_curvature_1_per_m", deflection.effective_curvature_1_per_m])
        deflection_sheet.append(["support_scheme", deflection.support_scheme])
        deflection_sheet.append(["k_m", deflection.k_m])
        deflection_sheet.append(["span_mm", deflection.span_mm])
        deflection_sheet.append(["a_mm", deflection.a_mm])
        deflection_sheet.append(["deflection_mm", deflection.deflection_mm])
        deflection_sheet.append(["limit_profile", deflection.limit.profile])
        deflection_sheet.append(["limit_mm", deflection.limit.limit_mm])
        deflection_sheet.append(["limit_rule_text", deflection.limit.rule_text])
        deflection_sheet.append(["requires_additional_data", deflection.limit.requires_additional_data])
        deflection_sheet.append(["warning_message", deflection.limit.warning_message])
        deflection_sheet.append(["is_within_limit", "OK" if deflection.is_within_limit else "NG"])
        deflection_sheet.append(["note", deflection.note])

        _add_deflection_theory_sheet(workbook, result, serviceability_report)

    return workbook


def build_results_workbook_bytes(
    section: SectionInput,
    catalog: MaterialCatalog,
    result: BendingResult,
    *,
    selected_point: CurvePoint | None = None,
    serviceability_report: ServiceabilityReport | None = None,
) -> bytes:
    buffer = BytesIO()
    workbook = build_results_workbook(
        section,
        catalog,
        result,
        selected_point=selected_point,
        serviceability_report=serviceability_report,
    )
    workbook.save(buffer)
    return buffer.getvalue()

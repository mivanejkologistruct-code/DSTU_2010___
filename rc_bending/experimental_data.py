from __future__ import annotations

from dataclasses import dataclass
from io import BytesIO
import re
from typing import Any

from openpyxl import Workbook, load_workbook
from openpyxl.utils import get_column_letter
import pandas as pd

MOMENT_COLUMN = "M, кН·м"

SUPPORTED_EXPERIMENTAL_SHEETS: dict[str, str] = {
    "M_f": "f, мм",
    "M_eps_c": "ε_c,top, 10^-5",
    "M_eps_s_top": "ε_s, 10^-5",
    "M_eps_s_bot": "ε_s, 10^-5",
}
REFERENCE_SUPPORTED_EXPERIMENTAL_SHEETS: dict[str, str] = {
    "M_f": "f, мм",
    "M_eps_c": "ε_c,top, 10^-5",
    "M_eps_s_bot": "ε_s, 10^-5",
}


@dataclass(frozen=True)
class ExperimentalSeries:
    sheet_name: str
    x_field: str
    data: pd.DataFrame


@dataclass(frozen=True)
class ExperimentalDataset:
    series: dict[str, ExperimentalSeries]
    ignored_sheets: tuple[str, ...] = ()

    def get_chart_df(self, sheet_name: str) -> pd.DataFrame | None:
        series = self.series.get(sheet_name)
        if series is None:
            return None
        return series.data.copy()


@dataclass(frozen=True)
class DetectedExperimentalBlock:
    target_graph: str
    sheet_name: str
    cell_range: str
    headers: tuple[str, str]
    preview_rows: tuple[tuple[object | None, object | None], ...]
    evidence_labels: tuple[str, ...]
    score: float


@dataclass(frozen=True)
class WorkbookInspectionResult:
    status: str
    dataset: ExperimentalDataset | None
    detected_blocks: tuple[DetectedExperimentalBlock, ...]
    missing_graphs: tuple[str, ...]
    conflicts: dict[str, tuple[str, ...]]
    message: str


@dataclass(frozen=True)
class _CandidateBlock:
    sheet_name: str
    row_start: int
    row_end: int
    col_start: int
    col_end: int
    headers: tuple[str, str]
    raw_rows: tuple[tuple[object | None, object | None], ...]
    evidence_labels: tuple[str, ...]
    target_scores: dict[str, float]

    @property
    def cell_range(self) -> str:
        return f"{get_column_letter(self.col_start)}{self.row_start}:{get_column_letter(self.col_end)}{self.row_end}"


def _build_template_workbook_bytes(supported_sheets: dict[str, str]) -> bytes:
    workbook = Workbook()
    first_sheet = workbook.active
    for sheet_index, (sheet_name, x_field) in enumerate(supported_sheets.items()):
        sheet = first_sheet if sheet_index == 0 else workbook.create_sheet(sheet_name)
        sheet.title = sheet_name
        sheet.append([x_field, MOMENT_COLUMN])

    buffer = BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def build_experimental_template_workbook_bytes() -> bytes:
    return _build_template_workbook_bytes(SUPPORTED_EXPERIMENTAL_SHEETS)


def build_reference_experimental_template_workbook_bytes() -> bytes:
    return _build_template_workbook_bytes(REFERENCE_SUPPORTED_EXPERIMENTAL_SHEETS)


def _load_numeric_sheet(df: pd.DataFrame, *, sheet_name: str, x_field: str) -> pd.DataFrame:
    normalized_columns = [str(column).strip() for column in df.columns]
    df = df.copy()
    df.columns = normalized_columns
    required_columns = [x_field, MOMENT_COLUMN]
    missing_columns = [column for column in required_columns if column not in df.columns]
    if missing_columns:
        raise ValueError(
            f"Аркуш {sheet_name} не містить обов'язкових колонок: {', '.join(missing_columns)}."
        )

    numeric_candidates = df.loc[:, required_columns].replace(r"^\s*$", pd.NA, regex=True)
    numeric_candidates = numeric_candidates.dropna(how="all").reset_index(drop=True)
    if numeric_candidates.empty or len(numeric_candidates.index) < 2:
        raise ValueError(f"Аркуш {sheet_name} має містити щонайменше 2 валідні точки.")

    numeric_df = numeric_candidates.apply(pd.to_numeric, errors="coerce")
    invalid_rows = numeric_df.isna().any(axis=1)
    if invalid_rows.any():
        invalid_index = int(invalid_rows[invalid_rows].index[0]) + 2
        raise ValueError(f"Аркуш {sheet_name} містить нечислові або неповні дані в рядку {invalid_index}.")

    return numeric_df


def _normalize_text(value: object | None) -> str:
    if value is None:
        return ""
    text = str(value).strip().lower()
    if not text:
        return ""
    substitutions = {
        "ε": " eps ",
        "_": " ",
        "-": " ",
        "–": " ",
        "—": " ",
        "/": " ",
        "\\": " ",
        "(": " ",
        ")": " ",
        ",": " ",
        ";": " ",
        ":": " ",
        ".": " ",
        "·": " ",
    }
    for source, replacement in substitutions.items():
        text = text.replace(source, replacement)
    text = re.sub(r"\s+", " ", text)
    return text.strip()


def _is_blank(value: object | None) -> bool:
    return value is None or (isinstance(value, str) and not value.strip())


def _is_numeric_like(value: object | None) -> bool:
    if _is_blank(value) or isinstance(value, bool):
        return False
    if isinstance(value, (int, float)):
        return True
    if isinstance(value, str):
        candidate = value.strip().replace(",", ".")
        if not candidate:
            return False
        try:
            float(candidate)
        except ValueError:
            return False
        return True
    return False


def _coerce_numeric(value: object | None) -> float | None:
    if _is_blank(value):
        return None
    if isinstance(value, bool):
        return None
    if isinstance(value, (int, float)):
        return float(value)
    if isinstance(value, str):
        candidate = value.strip().replace(",", ".")
        if not candidate:
            return None
        try:
            return float(candidate)
        except ValueError:
            return None
    return None


def _collect_evidence_labels(sheet, *, row_start: int, col_start: int, col_end: int) -> tuple[str, ...]:
    labels: list[str] = []
    row_candidates = [row_start - 1, row_start]
    col_candidates = [col_start - 1, col_start, col_end]
    for row_index in row_candidates:
        if row_index < 1:
            continue
        for col_index in col_candidates:
            if col_index < 1:
                continue
            value = sheet.cell(row_index, col_index).value
            if _is_blank(value):
                continue
            text = str(value).strip()
            if text not in labels:
                labels.append(text)
    return tuple(labels)


def _score_target_graph(*, headers: tuple[str, str], evidence_labels: tuple[str, ...], target_graph: str) -> float:
    header_text = " ".join(_normalize_text(value) for value in headers)
    evidence_text = " ".join(_normalize_text(value) for value in evidence_labels)
    if target_graph == "M_f":
        score = 0.0
        if "f" in header_text and "eps" not in header_text:
            score += 3.0
        if any(token in evidence_text for token in ("m f", "діаграма m f", "прогин", "deflection")):
            score += 2.0
        return score
    if target_graph == "M_eps_c":
        score = 0.0
        if "eps c" in header_text or "c top" in header_text:
            score += 3.0
        if any(token in evidence_text for token in ("бетон", "concrete")):
            score += 1.5
        return score
    if target_graph == "M_eps_s_top":
        score = 0.0
        if "eps s" in header_text:
            score += 2.0
        if any(token in evidence_text for token in ("верх", "top", "upper")):
            score += 2.5
        if any(token in evidence_text for token in ("арматур", "rebar", "steel")):
            score += 1.0
        if any(token in evidence_text for token in ("ниж", "bot", "bottom", "lower")):
            score -= 3.0
        return score
    if target_graph == "M_eps_s_bot":
        score = 0.0
        if "eps s" in header_text:
            score += 2.0
        if any(token in evidence_text for token in ("ниж", "bot", "bottom", "lower")):
            score += 2.5
        if any(token in evidence_text for token in ("арматур", "rebar", "steel")):
            score += 1.0
        if any(token in evidence_text for token in ("верх", "top", "upper")):
            score -= 3.0
        return score
    return 0.0


def _extract_candidate_blocks(sheet, *, supported_sheets: dict[str, str] = SUPPORTED_EXPERIMENTAL_SHEETS) -> tuple[_CandidateBlock, ...]:
    candidates: list[_CandidateBlock] = []
    max_row = sheet.max_row or 0
    max_column = sheet.max_column or 0
    for row_index in range(1, max_row + 1):
        for col_index in range(1, max_column):
            left_header = sheet.cell(row_index, col_index).value
            right_header = sheet.cell(row_index, col_index + 1).value
            if _is_blank(left_header) or _is_blank(right_header):
                continue
            if _is_numeric_like(left_header) and _is_numeric_like(right_header):
                continue

            raw_rows: list[tuple[object | None, object | None]] = []
            cursor = row_index + 1
            while cursor <= max_row:
                left_value = sheet.cell(cursor, col_index).value
                right_value = sheet.cell(cursor, col_index + 1).value
                if _is_blank(left_value) and _is_blank(right_value):
                    break
                raw_rows.append((left_value, right_value))
                cursor += 1
            if len(raw_rows) < 2:
                continue
            if not any(
                _coerce_numeric(left_value) is not None and _coerce_numeric(right_value) is not None
                for left_value, right_value in raw_rows
            ):
                continue

            headers = (str(left_header).strip(), str(right_header).strip())
            evidence_labels = _collect_evidence_labels(
                sheet,
                row_start=row_index,
                col_start=col_index,
                col_end=col_index + 1,
            )
            target_scores = {
                target_graph: _score_target_graph(headers=headers, evidence_labels=evidence_labels, target_graph=target_graph)
                for target_graph in supported_sheets
            }
            if max(target_scores.values(), default=0.0) <= 0.0:
                continue
            candidates.append(
                _CandidateBlock(
                    sheet_name=sheet.title,
                    row_start=row_index,
                    row_end=cursor - 1,
                    col_start=col_index,
                    col_end=col_index + 1,
                    headers=headers,
                    raw_rows=tuple(raw_rows),
                    evidence_labels=evidence_labels,
                    target_scores=target_scores,
                )
            )
    return tuple(candidates)


def _build_numeric_df_from_block(
    candidate: _CandidateBlock,
    *,
    target_graph: str,
    supported_sheets: dict[str, str] = SUPPORTED_EXPERIMENTAL_SHEETS,
) -> pd.DataFrame:
    x_field = supported_sheets[target_graph]
    raw_df = pd.DataFrame(candidate.raw_rows, columns=[x_field, MOMENT_COLUMN])
    numeric_candidates = raw_df.replace(r"^\s*$", pd.NA, regex=True)
    numeric_candidates = numeric_candidates.dropna(how="all").reset_index(drop=True)
    if numeric_candidates.empty or len(numeric_candidates.index) < 2:
        raise ValueError(
            f"Знайдений блок {candidate.cell_range} на аркуші {candidate.sheet_name} має містити щонайменше 2 валідні точки."
        )

    numeric_df = numeric_candidates.apply(pd.to_numeric, errors="coerce")
    invalid_rows = numeric_df.isna().any(axis=1)
    if invalid_rows.any():
        invalid_offset = int(invalid_rows[invalid_rows].index[0]) + 2
        invalid_row = candidate.row_start + invalid_offset
        raise ValueError(
            f"Знайдений блок {candidate.cell_range} на аркуші {candidate.sheet_name} містить нечислові або неповні дані в рядку {invalid_row}."
        )
    return numeric_df


def _candidate_to_detected_block(candidate: _CandidateBlock, *, target_graph: str) -> DetectedExperimentalBlock:
    return DetectedExperimentalBlock(
        target_graph=target_graph,
        sheet_name=candidate.sheet_name,
        cell_range=candidate.cell_range,
        headers=candidate.headers,
        preview_rows=tuple(candidate.raw_rows[:5]),
        evidence_labels=candidate.evidence_labels,
        score=float(candidate.target_scores.get(target_graph, 0.0)),
    )


def _inspect_single_sheet(
    sheet,
    *,
    supported_sheets: dict[str, str] = SUPPORTED_EXPERIMENTAL_SHEETS,
) -> WorkbookInspectionResult:
    candidates = _extract_candidate_blocks(sheet, supported_sheets=supported_sheets)
    matches_by_target: dict[str, list[_CandidateBlock]] = {target_graph: [] for target_graph in supported_sheets}
    detected_blocks: list[DetectedExperimentalBlock] = []

    for candidate in candidates:
        best_score = max(candidate.target_scores.values(), default=0.0)
        if best_score <= 0.0:
            continue
        matched_targets = [
            target_graph
            for target_graph, score in candidate.target_scores.items()
            if abs(score - best_score) <= 1e-9 and score > 0.0
        ]
        for target_graph in matched_targets:
            matches_by_target[target_graph].append(candidate)
            detected_blocks.append(_candidate_to_detected_block(candidate, target_graph=target_graph))

    conflicts = {
        target_graph: tuple(candidate.cell_range for candidate in matches)
        for target_graph, matches in matches_by_target.items()
        if len(matches) > 1
    }
    missing_graphs = tuple(
        target_graph for target_graph, matches in matches_by_target.items() if len(matches) == 0
    )

    if conflicts:
        conflict_targets = ", ".join(sorted(conflicts))
        return WorkbookInspectionResult(
            status="conflict",
            dataset=None,
            detected_blocks=tuple(detected_blocks),
            missing_graphs=missing_graphs,
            conflicts=conflicts,
            message=f"На аркуші {sheet.title} знайдено кілька кандидатів для графіків: {conflict_targets}.",
        )

    if missing_graphs:
        return WorkbookInspectionResult(
            status="incomplete",
            dataset=None,
            detected_blocks=tuple(detected_blocks),
            missing_graphs=missing_graphs,
            conflicts={},
            message=(
                f"На аркуші {sheet.title} не знайдено всі {len(supported_sheets)} графіки. "
                f"Відсутні: {', '.join(missing_graphs)}."
            ),
        )

    series: dict[str, ExperimentalSeries] = {}
    ready_blocks: list[DetectedExperimentalBlock] = []
    for target_graph, matches in matches_by_target.items():
        candidate = matches[0]
        numeric_df = _build_numeric_df_from_block(candidate, target_graph=target_graph, supported_sheets=supported_sheets)
        series[target_graph] = ExperimentalSeries(
            sheet_name=target_graph,
            x_field=supported_sheets[target_graph],
            data=numeric_df,
        )
        ready_blocks.append(_candidate_to_detected_block(candidate, target_graph=target_graph))

    return WorkbookInspectionResult(
        status="ready",
        dataset=ExperimentalDataset(series=series, ignored_sheets=()),
        detected_blocks=tuple(ready_blocks),
        missing_graphs=(),
        conflicts={},
        message=f"На аркуші {sheet.title} розпізнано всі {len(supported_sheets)} графіки.",
    )


def inspect_experimental_workbook(
    workbook_bytes: bytes,
    *,
    supported_sheets: dict[str, str] = SUPPORTED_EXPERIMENTAL_SHEETS,
) -> WorkbookInspectionResult:
    workbook = load_workbook(BytesIO(workbook_bytes), data_only=True)
    sheet_results = [_inspect_single_sheet(sheet, supported_sheets=supported_sheets) for sheet in workbook.worksheets]
    ready_results = [result for result in sheet_results if result.status == "ready"]

    if len(ready_results) == 1:
        return ready_results[0]
    if len(ready_results) > 1:
        all_blocks = tuple(block for result in ready_results for block in result.detected_blocks)
        return WorkbookInspectionResult(
            status="conflict",
            dataset=None,
            detected_blocks=all_blocks,
            missing_graphs=(),
            conflicts={"workbook": tuple(block.sheet_name for block in all_blocks)},
            message=f"У книзі знайдено більше одного аркуша з повним набором із {len(supported_sheets)} графіків.",
        )

    if not sheet_results:
        return WorkbookInspectionResult(
            status="incomplete",
            dataset=None,
            detected_blocks=(),
            missing_graphs=tuple(supported_sheets),
            conflicts={},
            message="Книга Excel не містить жодного аркуша для аналізу.",
        )

    def _result_rank(result: WorkbookInspectionResult) -> tuple[int, int, int]:
        return (
            len(result.detected_blocks),
            -len(result.conflicts),
            -len(result.missing_graphs),
        )

    best_result = max(sheet_results, key=_result_rank)
    if best_result.detected_blocks:
        return best_result
    return WorkbookInspectionResult(
        status="incomplete",
        dataset=None,
        detected_blocks=(),
        missing_graphs=tuple(supported_sheets),
        conflicts={},
        message="Не вдалося розпізнати жодної таблиці експериментальних графіків у завантаженому файлі.",
    )


def inspect_reference_experimental_workbook(workbook_bytes: bytes) -> WorkbookInspectionResult:
    return inspect_experimental_workbook(workbook_bytes, supported_sheets=REFERENCE_SUPPORTED_EXPERIMENTAL_SHEETS)


def _load_legacy_supported_dataset(
    workbook_bytes: bytes,
    *,
    supported_sheets: dict[str, str],
) -> ExperimentalDataset | None:
    excel_file = pd.ExcelFile(BytesIO(workbook_bytes))
    supported_sheet_names = [sheet_name for sheet_name in excel_file.sheet_names if sheet_name in supported_sheets]
    if not supported_sheet_names:
        return None

    series: dict[str, ExperimentalSeries] = {}
    ignored_sheets = tuple(
        sheet_name for sheet_name in excel_file.sheet_names if sheet_name not in supported_sheets
    )

    for sheet_name, x_field in supported_sheets.items():
        if sheet_name not in excel_file.sheet_names:
            continue
        parsed_df = excel_file.parse(sheet_name=sheet_name)
        numeric_df = _load_numeric_sheet(parsed_df, sheet_name=sheet_name, x_field=x_field)
        series[sheet_name] = ExperimentalSeries(sheet_name=sheet_name, x_field=x_field, data=numeric_df)

    if not series:
        return None

    return ExperimentalDataset(series=series, ignored_sheets=ignored_sheets)


def load_experimental_dataset(workbook_bytes: bytes) -> ExperimentalDataset:
    legacy_dataset = _load_legacy_supported_dataset(workbook_bytes, supported_sheets=SUPPORTED_EXPERIMENTAL_SHEETS)
    if legacy_dataset is not None:
        return legacy_dataset

    inspection_result = inspect_experimental_workbook(workbook_bytes)
    if inspection_result.dataset is None:
        if not inspection_result.detected_blocks:
            raise ValueError("Файл не містить жодного підтримуваного аркуша з експериментальними кривими.")
        raise ValueError(inspection_result.message)
    return inspection_result.dataset


def load_reference_experimental_dataset(workbook_bytes: bytes) -> ExperimentalDataset:
    legacy_dataset = _load_legacy_supported_dataset(
        workbook_bytes,
        supported_sheets=REFERENCE_SUPPORTED_EXPERIMENTAL_SHEETS,
    )
    if legacy_dataset is not None:
        return legacy_dataset

    inspection_result = inspect_reference_experimental_workbook(workbook_bytes)
    if inspection_result.dataset is None:
        if not inspection_result.detected_blocks:
            raise ValueError("Файл не містить жодного підтримуваного аркуша з експериментальними кривими еталонної плити.")
        raise ValueError(inspection_result.message)
    return inspection_result.dataset


def _sort_theory_df(theory_df: pd.DataFrame, *, x_field: str) -> pd.DataFrame:
    sort_fields = [x_field]
    if "Крок" in theory_df.columns:
        sort_fields.append("Крок")
    return theory_df.sort_values(by=sort_fields, kind="mergesort").reset_index(drop=True)


def _interpolate_theory_moment(theory_df: pd.DataFrame, *, x_field: str, target_x: float) -> float | None:
    sorted_df = _sort_theory_df(theory_df, x_field=x_field)
    x_values = [float(value) for value in sorted_df[x_field].tolist()]
    moment_values = [float(value) for value in sorted_df[MOMENT_COLUMN].tolist()]

    tolerance = 1e-9
    if target_x < x_values[0] - tolerance or target_x > x_values[-1] + tolerance:
        return None

    for x_value, moment in zip(x_values, moment_values, strict=True):
        if abs(x_value - target_x) <= tolerance:
            return moment

    for index in range(1, len(x_values)):
        left_x = x_values[index - 1]
        right_x = x_values[index]
        if not (left_x <= target_x <= right_x):
            continue
        left_moment = moment_values[index - 1]
        right_moment = moment_values[index]
        if abs(right_x - left_x) <= tolerance:
            return right_moment
        ratio = (target_x - left_x) / (right_x - left_x)
        return left_moment + ratio * (right_moment - left_moment)

    return None


def build_comparison_table(theory_df: pd.DataFrame, experimental_series: ExperimentalSeries) -> pd.DataFrame:
    rows: list[dict[str, float | str | None]] = []
    tolerance = 1e-9

    for _, row in experimental_series.data.iterrows():
        x_value = float(row[experimental_series.x_field])
        experimental_moment = float(row[MOMENT_COLUMN])
        theory_moment = _interpolate_theory_moment(theory_df, x_field=experimental_series.x_field, target_x=x_value)

        if theory_moment is None:
            rows.append(
                {
                    "x_exp": x_value,
                    "M_exp, кН·м": experimental_moment,
                    "M_theory_interp, кН·м": None,
                    "ΔM, кН·м": None,
                    "ΔM, %": None,
                    "Статус": "Поза діапазоном теорії",
                }
            )
            continue

        delta_moment = experimental_moment - theory_moment
        delta_percent = None if abs(theory_moment) <= tolerance else 100.0 * delta_moment / theory_moment
        rows.append(
            {
                "x_exp": x_value,
                "M_exp, кН·м": experimental_moment,
                "M_theory_interp, кН·м": theory_moment,
                "ΔM, кН·м": delta_moment,
                "ΔM, %": delta_percent,
                "Статус": "OK",
            }
        )

    return pd.DataFrame(rows)

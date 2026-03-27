from dataclasses import replace
from io import BytesIO
import json
from pathlib import Path

from openpyxl import Workbook, load_workbook
import pandas as pd
import pytest
from streamlit.testing.v1 import AppTest

from rc_bending.export import build_results_workbook_bytes
from rc_bending.materials import load_material_catalog
from rc_bending.models import ChartLimitAnnotation, ConcreteLayerInput, CurvePoint, RebarLayerInput, SectionInput, ServiceabilityInput
from rc_bending.solver import build_layer_force_table, build_strain_profile_for_point, solve_bending_capacity
from rc_bending.ui_helpers import (
    add_rebar_layer,
    build_section_input_from_draft,
    default_draft_inputs,
    derive_draft_geometry,
    remove_rebar_layer,
    validate_draft_inputs,
)


def make_case() -> SectionInput:
    return SectionInput(
        section_height_mm=500.0,
        concrete_layers=(
            ConcreteLayerInput(width_mm=500.0, height_mm=200.0, concrete_class="C30/35"),
            ConcreteLayerInput(width_mm=500.0, height_mm=300.0, concrete_class="C20/25"),
        ),
        rebar_layers=(
            RebarLayerInput(z_mm=30.0, area_mm2=7 * 615.8, steel_class="A400C"),
            RebarLayerInput(z_mm=470.0, area_mm2=7 * 615.8, steel_class="A400C"),
        ),
    )


def _find_widget_by_label(elements, label: str):
    for element in elements:
        if getattr(element, "label", None) == label:
            return element
    raise AssertionError(f"Widget with label '{label}' was not found.")


def _find_metric(app: AppTest, label: str):
    return _find_widget_by_label(app.metric, label)


def _has_metric(app: AppTest, label: str) -> bool:
    return any(getattr(metric, "label", None) == label for metric in app.metric)


def _find_dataframe(app: AppTest, columns: list[str]):
    for dataframe in app.dataframe:
        if list(dataframe.value.columns) == columns:
            return dataframe.value
    raise AssertionError(f"DataFrame with columns {columns!r} was not found.")


def _has_dataframe(app: AppTest, columns: list[str]) -> bool:
    return any(list(dataframe.value.columns) == columns for dataframe in app.dataframe)


def _find_markdown_containing(app: AppTest, substring: str):
    for markdown in app.markdown:
        if substring in markdown.value:
            return markdown.value
    raise AssertionError(f"Markdown containing {substring!r} was not found.")


def _has_markdown_containing(app: AppTest, substring: str) -> bool:
    return any(substring in markdown.value for markdown in app.markdown)


def _count_markdown_containing(app: AppTest, substring: str) -> int:
    return sum(1 for markdown in app.markdown if substring in markdown.value)


def _find_tab(app: AppTest, label: str):
    for tab in app.tabs:
        if getattr(tab, "label", None) == label:
            return tab
    raise AssertionError(f"Tab with label {label!r} was not found.")


def _workspace_button_label(workspace: str) -> str:
    return {
        "section": "Переріз",
        "serviceability": "II ГГС",
        "experimental": "Експеримент",
    }[workspace]


def _activate_workspace(app: AppTest, workspace: str):
    _find_widget_by_label(app.button, _workspace_button_label(workspace)).click()
    app.run(timeout=10)
    return app.main


def _find_markdown_containing_in_node(node, substring: str):
    for markdown in node.markdown:
        if substring in markdown.value:
            return markdown.value
    raise AssertionError(f"Markdown containing {substring!r} was not found in node.")


def _has_markdown_containing_in_node(node, substring: str) -> bool:
    return any(substring in markdown.value for markdown in node.markdown)


def _find_widget_by_label_in_node(node, collection_name: str, label: str):
    return _find_widget_by_label(getattr(node, collection_name), label)


def _block_and_child_index_with_markdown_in_node(node, substring: str) -> tuple[int, int]:
    for index, element in node.children.items():
        if type(element).__name__ == "Markdown" and substring in element.value:
            return index, 0
        if type(element).__name__ != "Block":
            continue
        descendant_order = 0
        for child in _iter_descendants(element):
            if type(child).__name__ == "Markdown" and substring in child.value:
                return index, descendant_order
            descendant_order += 1
    raise AssertionError(f"Block markdown containing {substring!r} was not found in node.")


def _block_and_child_index_with_subheader_in_node(node, title: str) -> tuple[int, int]:
    for index, element in node.children.items():
        if type(element).__name__ != "Block":
            continue
        descendant_order = 0
        for child in _iter_descendants(element):
            if getattr(child, "type", None) == "subheader" and getattr(child, "value", None) == title:
                return index, descendant_order
            descendant_order += 1
    raise AssertionError(f"Block subheader {title!r} was not found in node.")


def _iter_descendants(node):
    children = getattr(node, "children", None)
    if not isinstance(children, dict):
        return
    for child in children.values():
        yield child
        yield from _iter_descendants(child)


def _top_level_index_markdown_containing(app: AppTest, substring: str) -> int:
    for index, element in app.main.children.items():
        if type(element).__name__ == "Markdown" and substring in element.value:
            return index
    raise AssertionError(f"Top-level markdown containing {substring!r} was not found.")


def _top_level_index_widget(app: AppTest, *, widget_type: str, label: str) -> int:
    for index, element in app.main.children.items():
        if getattr(element, "type", None) == widget_type and getattr(element, "label", None) == label:
            return index
    raise AssertionError(f"Top-level widget {widget_type!r} with label {label!r} was not found.")


def _top_level_block_index_with_subheader(app: AppTest, title: str) -> int:
    for index, element in app.main.children.items():
        if type(element).__name__ != "Block":
            continue
        if any(getattr(child, "value", None) == title for child in element.subheader):
            return index
    raise AssertionError(f"Top-level block with subheader {title!r} was not found.")


def _top_level_block_index_with_markdown(app: AppTest, substring: str) -> int:
    for index, element in app.main.children.items():
        if type(element).__name__ == "Markdown" and substring in getattr(element, "value", ""):
            return index
        if type(element).__name__ != "Block":
            continue
        if any(substring in getattr(child, "value", "") for child in element.markdown):
            return index
    raise AssertionError(f"Top-level block with markdown containing {substring!r} was not found.")


def _top_level_block_index_by_type(app: AppTest, block_type: str) -> int:
    for index, element in app.main.children.items():
        if type(element).__name__ == "Block" and getattr(element, "type", None) == block_type:
            return index
    raise AssertionError(f"Top-level block with type {block_type!r} was not found.")


def _vega_spec(element) -> dict:
    return json.loads(element.proto.spec)


def _build_experimental_workbook_bytes(sheet_rows: dict[str, list[list[object]]]) -> bytes:
    workbook = Workbook()
    first_sheet = workbook.active
    first_title = next(iter(sheet_rows))
    first_sheet.title = first_title

    for index, (sheet_name, rows) in enumerate(sheet_rows.items()):
        sheet = first_sheet if index == 0 else workbook.create_sheet(sheet_name)
        for row in rows:
            sheet.append(row)

    buffer = BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def _build_positioned_experimental_workbook_bytes(
    sheet_cells: dict[str, object],
    *,
    sheet_name: str = "Experiment",
) -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = sheet_name

    for cell_ref, value in sheet_cells.items():
        sheet[cell_ref] = value

    buffer = BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def _first_vega_spec_in_block_with_subheader(app: AppTest, title: str) -> dict:
    block_index = _top_level_block_index_with_subheader(app, title)
    block = app.main.children[block_index]
    for child in block.children.values():
        if getattr(child, "type", None) == "vega_lite_chart":
            return _vega_spec(child)
    raise AssertionError(f"Vega chart inside block {title!r} was not found.")


def _first_top_level_vega_spec_after_subheader(app: AppTest, title: str) -> dict:
    subheader_found = False
    for element in app.main.children.values():
        if getattr(element, "type", None) == "subheader" and getattr(element, "value", None) == title:
            subheader_found = True
            continue
        if not subheader_found:
            continue
        if getattr(element, "type", None) == "vega_lite_chart":
            return _vega_spec(element)
        if getattr(element, "type", None) == "subheader":
            break
    raise AssertionError(f"Top-level vega chart after subheader {title!r} was not found.")


def _find_block_with_subheader_in_node(node, title: str):
    matching_blocks = []
    for child in _iter_descendants(node):
        if type(child).__name__ != "Block":
            continue
        titles = [getattr(subheader, "value", None) for subheader in child.subheader]
        if any(candidate == title for candidate in titles):
            matching_blocks.append((len(titles), sum(1 for _ in _iter_descendants(child)), child))
    if matching_blocks:
        matching_blocks.sort(key=lambda item: (item[0], item[1]))
        return matching_blocks[0][2]
    raise AssertionError(f"Block with subheader {title!r} was not found in node.")


def _first_vega_spec_in_node(node) -> dict:
    for child in _iter_descendants(node):
        if getattr(child, "type", None) == "vega_lite_chart":
            return _vega_spec(child)
    raise AssertionError("Vega chart was not found in node.")


def _first_vega_spec_in_node_block_with_subheader(node, title: str) -> dict:
    block = _find_block_with_subheader_in_node(node, title)
    return _first_vega_spec_in_node(block)


def test_export_workbook_contains_required_sheets_and_charts(tmp_path):
    section = make_case()
    catalog = load_material_catalog()
    result = solve_bending_capacity(section, catalog)

    workbook_bytes = build_results_workbook_bytes(section, catalog, result)
    output_path = tmp_path / "result.xlsx"
    output_path.write_bytes(workbook_bytes)
    workbook = load_workbook(output_path)

    assert workbook.sheetnames == [
        "Inputs",
        "Materials",
        "MomentCurvature",
        "ConcreteStrainProfile",
        "IntermediateIterations",
        "LayerForces",
        "ConcreteMomentStrainTheory",
        "TopRebarMomentStrainTheory",
        "BottomRebarMomentStrainTheory",
    ]
    assert workbook["MomentCurvature"]["A2"].value == 1
    assert workbook["MomentCurvature"]["D2"].value == 0
    assert workbook["MomentCurvature"]["D3"].value > 0
    assert workbook["ConcreteStrainProfile"]["A2"].value == 0
    assert workbook["IntermediateIterations"].max_row > 2
    assert workbook["LayerForces"].max_row > 2
    assert workbook["ConcreteMomentStrainTheory"]["A1"].value == "Крок"
    assert workbook["ConcreteMomentStrainTheory"]["B1"].value == "M, кН·м"
    assert workbook["ConcreteMomentStrainTheory"]["C1"].value == "ε_c,top, 10^-5"
    assert workbook["TopRebarMomentStrainTheory"]["C1"].value == "ε_s,top, 10^-5"
    assert workbook["BottomRebarMomentStrainTheory"]["C1"].value == "ε_s,bot, 10^-5"
    assert len(workbook["MomentCurvature"]._charts) == 1
    assert len(workbook["ConcreteStrainProfile"]._charts) == 1
    assert len(workbook["ConcreteMomentStrainTheory"]._charts) == 1
    assert len(workbook["TopRebarMomentStrainTheory"]._charts) == 1
    assert len(workbook["BottomRebarMomentStrainTheory"]._charts) == 1


def test_export_can_follow_a_selected_curve_point(tmp_path):
    section = make_case()
    catalog = load_material_catalog()
    result = solve_bending_capacity(section, catalog)
    selected_point = result.curve_points[0]
    expected_profile = build_strain_profile_for_point(section, selected_point)

    workbook_bytes = build_results_workbook_bytes(section, catalog, result, selected_point=selected_point)
    output_path = tmp_path / "selected_result.xlsx"
    output_path.write_bytes(workbook_bytes)
    workbook = load_workbook(output_path)

    assert workbook["ConcreteStrainProfile"]["B2"].value == expected_profile[0].strain


def test_export_workbook_includes_serviceability_sheets_when_report_is_provided(tmp_path):
    from rc_bending.serviceability import build_serviceability_report

    catalog = load_material_catalog()
    draft = default_draft_inputs()
    draft["serviceability"]["span_mm"] = 6000.0
    section = build_section_input_from_draft(draft, catalog)
    result = solve_bending_capacity(section, catalog)
    selected_point = result.curve_points[-1]
    report = build_serviceability_report(section, draft, selected_point, catalog)

    workbook_bytes = build_results_workbook_bytes(
        section,
        catalog,
        result,
        selected_point=selected_point,
        serviceability_report=report,
    )
    output_path = tmp_path / "serviceability_result.xlsx"
    output_path.write_bytes(workbook_bytes)
    workbook = load_workbook(output_path)

    assert "ServiceabilityInputs" in workbook.sheetnames
    assert "CrackWidthCheck" in workbook.sheetnames
    assert "DeflectionCheck" in workbook.sheetnames
    assert "DeflectionCurveTheory" in workbook.sheetnames
    assert workbook["DeflectionCurveTheory"]["A1"].value == "Крок"
    assert workbook["DeflectionCurveTheory"]["B1"].value == "M, кН·м"
    assert workbook["DeflectionCurveTheory"]["C1"].value == "f, мм"
    assert len(workbook["DeflectionCurveTheory"]._charts) == 1


def test_export_workbook_uses_single_rebar_theory_sheet_for_single_layer_section(tmp_path):
    catalog = load_material_catalog()
    draft = default_draft_inputs()
    draft["has_strengthening_layer"] = False
    draft["rebar_layers"] = [draft["rebar_layers"][1]]
    section = build_section_input_from_draft(draft, catalog)
    result = solve_bending_capacity(section, catalog)

    workbook_bytes = build_results_workbook_bytes(section, catalog, result)
    output_path = tmp_path / "single-layer-result.xlsx"
    output_path.write_bytes(workbook_bytes)
    workbook = load_workbook(output_path)

    assert "RebarMomentStrainTheory" in workbook.sheetnames
    assert "TopRebarMomentStrainTheory" not in workbook.sheetnames
    assert "BottomRebarMomentStrainTheory" not in workbook.sheetnames
    assert workbook["RebarMomentStrainTheory"]["A1"].value == "Крок"
    assert workbook["RebarMomentStrainTheory"]["B1"].value == "M, кН·м"
    assert workbook["RebarMomentStrainTheory"]["C1"].value == "ε_s, 10^-5"
    assert len(workbook["RebarMomentStrainTheory"]._charts) == 1


def test_ui_helpers_can_add_and_remove_rebar_rows():
    rows = [
        {"id": "rebar_1", "face": "top", "distance_mm": 30.0, "bar_count": 7, "diameter_mm": 28, "steel_class": "A400C"},
        {"id": "rebar_2", "face": "bottom", "distance_mm": 30.0, "bar_count": 7, "diameter_mm": 28, "steel_class": "A400C"},
    ]

    rows = add_rebar_layer(rows)
    assert len(rows) == 3
    assert rows[-1]["bar_count"] == 1
    assert rows[-1]["face"] == "Верхня"

    rows = remove_rebar_layer(rows, 1)
    assert len(rows) == 2
    assert rows[1]["id"] != "rebar_2"


def test_current_point_cards_html_contains_named_value_slots():
    from streamlit_app import _build_current_point_cards_html

    catalog = load_material_catalog()
    result = solve_bending_capacity(make_case(), catalog)

    html = _build_current_point_cards_html(result.peak_point)

    assert 'data-role="point-chip-grid"' in html
    assert html.count('data-role="point-chip"') == 5
    assert 'data-role="point-value-step"' in html
    assert 'data-role="point-value-moment"' in html
    assert 'data-role="point-value-curvature"' in html
    assert 'data-role="point-value-neutral-axis"' in html
    assert 'data-role="point-value-residual"' in html


def test_custom_css_uses_wider_content_width_for_drawing_readability():
    from streamlit_app import _build_custom_css

    css = _build_custom_css()

    assert "max-width: 1480px" in css
    assert 'div[data-baseweb="input"] > div' in css
    assert 'div[data-baseweb="select"] > div' in css
    assert "border: 1.5px solid" in css
    assert "box-shadow:" in css
    assert ":focus-within > div" in css


def test_force_cards_html_contains_one_card_per_layer():
    from streamlit_app import _build_force_cards_html

    catalog = load_material_catalog()
    section = make_case()
    result = solve_bending_capacity(section, catalog)
    layer_force_rows = build_layer_force_table(section, catalog, result.peak_point)

    html = _build_force_cards_html(layer_force_rows)

    assert 'data-role="force-card-grid"' in html
    assert html.count('data-role="force-card"') == len(layer_force_rows)
    assert "B1" in html
    assert "A2" in html
    assert "N, кН" in html
    assert "σ, МПа" in html


def test_drawing_showcase_html_wraps_svg_with_summary_chips():
    from streamlit_app import _build_drawing_showcase_html
    from rc_bending.section_drawing import build_section_drawing_svg
    from rc_bending.ui_helpers import default_draft_inputs, derive_draft_geometry

    catalog = load_material_catalog()
    draft = default_draft_inputs()
    derived = derive_draft_geometry(draft)
    section = build_section_input_from_draft(draft, catalog)
    result = solve_bending_capacity(section, catalog)

    html = _build_drawing_showcase_html(
        derived,
        build_section_drawing_svg(derived, selected_point=result.peak_point),
        selected_point=result.peak_point,
        note="Тестова примітка.",
    )

    assert 'data-role="drawing-showcase"' in html
    assert 'data-role="drawing-showcase-chip"' in html
    assert "Переріз та епюри форм рівноваги" in html
    assert "Тестова примітка." in html
    assert "120.0 мм" in html


def test_author_cards_html_renders_profiles():
    from streamlit_app import _build_hero_banner_html

    html = _build_hero_banner_html()

    assert 'data-role="hero-author-list"' in html
    assert html.count('data-role="hero-author-item"') == 2
    assert "Іванейко М.М." in html
    assert "Іванейко В.М." in html
    assert "Автор проєкту." in html


def test_draft_validation_requires_mode_specific_rebar_count():
    catalog = load_material_catalog()
    one_layer_draft = default_draft_inputs()
    one_layer_draft["rebar_layers"] = one_layer_draft["rebar_layers"][:1]

    three_layer_draft = default_draft_inputs()
    three_layer_draft["rebar_layers"] = add_rebar_layer(list(three_layer_draft["rebar_layers"]))

    one_layer_errors = validate_draft_inputs(one_layer_draft, catalog)
    three_layer_errors = validate_draft_inputs(three_layer_draft, catalog)

    assert any("Рівно два шари арматури" in error for error in one_layer_errors)
    assert any("Рівно два шари арматури" in error for error in three_layer_errors)

    reference_draft = default_draft_inputs()
    reference_draft["has_strengthening_layer"] = False
    reference_draft["rebar_layers"] = [reference_draft["rebar_layers"][1]]

    assert validate_draft_inputs(reference_draft, catalog) == []


def test_build_section_input_from_draft_rejects_non_normative_rebar_count():
    catalog = load_material_catalog()
    draft = default_draft_inputs()
    draft["rebar_layers"] = add_rebar_layer(list(draft["rebar_layers"]))

    with pytest.raises(ValueError, match="One or two rebar layers are required"):
        build_section_input_from_draft(draft, catalog)


def test_build_section_input_from_draft_normalizes_reference_slab_to_base_material():
    catalog = load_material_catalog()
    draft = default_draft_inputs()
    draft["has_strengthening_layer"] = False
    draft["concrete_layers"][0]["height_mm"] = 20.0
    draft["concrete_layers"][0]["concrete_class"] = "C40/50"
    draft["concrete_layers"][1]["concrete_class"] = "C25/30"
    draft["rebar_layers"] = [draft["rebar_layers"][1]]

    assert validate_draft_inputs(draft, catalog) == []

    section = build_section_input_from_draft(draft, catalog)

    assert section.concrete_layers[0].height_mm == pytest.approx(float(draft["section_height_mm"]))
    assert section.concrete_layers[0].concrete_class == "C25/30"
    assert section.concrete_layers[1].height_mm == pytest.approx(0.0)
    assert section.concrete_layers[1].concrete_class == "C25/30"
    assert len(section.rebar_layers) == 1
    assert section.rebar_layers[0].z_mm == pytest.approx(100.0)


def test_reference_slab_preview_uses_single_active_concrete_layer():
    from streamlit_app import _build_concrete_preview_df, _build_drawing_showcase_html

    draft = default_draft_inputs()
    draft["has_strengthening_layer"] = False
    draft["concrete_layers"][0]["height_mm"] = 20.0
    draft["concrete_layers"][0]["concrete_class"] = "C40/50"
    draft["concrete_layers"][1]["concrete_class"] = "C25/30"
    draft["rebar_layers"] = [draft["rebar_layers"][1]]

    derived = derive_draft_geometry(draft)
    preview_df = _build_concrete_preview_df(derived)
    showcase_html = _build_drawing_showcase_html(derived, "<svg></svg>", note="Тестова примітка.")

    assert preview_df["Шар"].tolist() == ["B1"]
    assert preview_df["h, мм"].tolist() == [pytest.approx(float(draft["section_height_mm"]))]
    assert preview_df["Клас"].tolist() == ["C25/30"]
    assert "C40/50 / C25/30" not in showcase_html
    assert "C25/30" in showcase_html


def test_reference_slab_preview_uses_single_active_rebar_layer():
    from streamlit_app import _build_rebar_preview_df

    draft = default_draft_inputs()
    draft["has_strengthening_layer"] = False
    draft["rebar_layers"] = [draft["rebar_layers"][1]]

    derived = derive_draft_geometry(draft)
    preview_df = _build_rebar_preview_df(derived)

    assert preview_df["Шар"].tolist() == ["A1"]
    assert preview_df["Грань"].tolist() == ["Нижня"]
    assert preview_df["z, мм"].tolist() == [pytest.approx(100.0)]


def test_concrete_and_rebar_chart_data_use_1e5_strain_scale():
    from streamlit_app import (
        _build_concrete_moment_strain_df,
        _build_concrete_strain_chart_df,
        _build_rebar_moment_strain_df,
        _pick_extreme_rebar_layers,
    )

    catalog = load_material_catalog()
    section = make_case()
    result = solve_bending_capacity(section, catalog)
    top_rebar, bottom_rebar = _pick_extreme_rebar_layers(section)

    concrete_moment_df = _build_concrete_moment_strain_df(result)
    concrete_df = _build_concrete_strain_chart_df(build_strain_profile_for_point(section, result.peak_point))
    top_df = _build_rebar_moment_strain_df(section, result, rebar_index=top_rebar[0])
    bottom_df = _build_rebar_moment_strain_df(section, result, rebar_index=bottom_rebar[0])

    first_point = result.curve_points[0]
    top_expected = (
        first_point.top_strain
        + (first_point.bottom_strain - first_point.top_strain) * top_rebar[1].z_mm / section.section_height_mm
    ) * 100000.0
    bottom_expected_raw = (
        first_point.top_strain
        + (first_point.bottom_strain - first_point.top_strain) * bottom_rebar[1].z_mm / section.section_height_mm
    ) * 100000.0
    bottom_expected = -bottom_expected_raw

    assert list(concrete_moment_df.columns) == ["Крок", "M, кН·м", "ε_c,top, 10^-5"]
    assert list(concrete_df.columns) == ["z, мм", "ε_c, 10^-5"]
    assert list(top_df.columns) == ["Крок", "M, кН·м", "ε_s, 10^-5"]
    assert list(bottom_df.columns) == ["Крок", "M, кН·м", "ε_s, 10^-5"]
    assert concrete_moment_df.iloc[0]["ε_c,top, 10^-5"] == pytest.approx(first_point.top_strain * 100000.0)
    assert concrete_df.iloc[0]["ε_c, 10^-5"] == pytest.approx(result.peak_point.top_strain * 100000.0)
    assert any(value == pytest.approx(top_expected) for value in top_df["ε_s, 10^-5"].tolist())
    assert any(value == pytest.approx(bottom_expected) for value in bottom_df["ε_s, 10^-5"].tolist())
    assert concrete_moment_df["ε_c,top, 10^-5"].is_monotonic_increasing
    assert top_df["ε_s, 10^-5"].is_monotonic_increasing
    assert bottom_df["ε_s, 10^-5"].is_monotonic_increasing


def test_build_limit_annotation_returns_exact_match_when_limit_hits_curve_point():
    from streamlit_app import _build_limit_annotation

    points = (
        CurvePoint(1, 0.0, 0.0, 0.0, 500.0, 0.0, 0.0, "whole_compression"),
        CurvePoint(2, 0.001, -0.010, 2.0, 40.0, 0.0, 120.0, "bending_with_tension"),
    )

    annotation = _build_limit_annotation(
        points,
        x_values=[0.0, 100.0],
        target_strain=100.0,
        label="ε_limit",
    )

    assert annotation.label == "ε_limit"
    assert annotation.target_strain == pytest.approx(100.0)
    assert annotation.moment_kNm == pytest.approx(120.0)
    assert annotation.within_chart_range is True


def test_build_limit_annotation_interpolates_moment_between_curve_points():
    from streamlit_app import _build_limit_annotation

    points = (
        CurvePoint(1, 0.0, 0.0, 0.0, 500.0, 0.0, 100.0, "whole_compression"),
        CurvePoint(2, 0.001, -0.010, 2.0, 40.0, 0.0, 200.0, "bending_with_tension"),
    )

    annotation = _build_limit_annotation(
        points,
        x_values=[10.0, 30.0],
        target_strain=20.0,
        label="ε_limit",
    )

    assert annotation.moment_kNm == pytest.approx(150.0)
    assert annotation.within_chart_range is True


def test_build_limit_annotation_marks_out_of_range_limit_without_moment():
    from streamlit_app import _build_limit_annotation

    points = (
        CurvePoint(1, 0.0, 0.0, 0.0, 500.0, 0.0, 100.0, "whole_compression"),
        CurvePoint(2, 0.001, -0.010, 2.0, 40.0, 0.0, 200.0, "bending_with_tension"),
    )

    annotation = _build_limit_annotation(
        points,
        x_values=[10.0, 30.0],
        target_strain=45.0,
        label="ε_limit",
    )

    assert annotation.moment_kNm is None
    assert annotation.within_chart_range is False


def test_concrete_limit_annotation_uses_normative_limit_as_primary_boundary():
    from streamlit_app import _build_concrete_limit_annotation

    catalog = load_material_catalog()
    section = build_section_input_from_draft(default_draft_inputs(), catalog)
    result = solve_bending_capacity(section, catalog)

    annotation = _build_concrete_limit_annotation(section, catalog, result)

    assert annotation.label == "ε_cu1,ck"
    assert annotation.target_strain == pytest.approx(263.0)
    assert annotation.moment_kNm == pytest.approx(result.curve_points[-1].moment_kNm)
    assert annotation.within_chart_range is True
    assert annotation.secondary_label is None
    assert annotation.secondary_strain is None


def test_rebar_limit_annotation_uses_display_limit_for_top_negative_and_bottom_positive_branches():
    from streamlit_app import _build_rebar_limit_annotation, _pick_extreme_rebar_layers

    catalog = load_material_catalog()
    section = build_section_input_from_draft(default_draft_inputs(), catalog)
    result = solve_bending_capacity(section, catalog)
    top_rebar, bottom_rebar = _pick_extreme_rebar_layers(section)

    top_annotation = _build_rebar_limit_annotation(section, catalog, result, rebar_index=top_rebar[0])
    bottom_annotation = _build_rebar_limit_annotation(section, catalog, result, rebar_index=bottom_rebar[0])

    assert top_annotation.label == "ε_yk"
    assert top_annotation.target_strain == pytest.approx(-281.0)
    assert top_annotation.secondary_label == "ε_ud"
    assert top_annotation.secondary_strain == pytest.approx(-2000.0)
    assert bottom_annotation.label == "ε_yk"
    assert bottom_annotation.target_strain == pytest.approx(281.0)
    assert bottom_annotation.secondary_label == "ε_ud"
    assert bottom_annotation.secondary_strain == pytest.approx(2000.0)


def test_last_curve_point_annotations_use_epsilon_max_for_final_curve_point():
    from streamlit_app import (
        _build_concrete_last_point_annotation,
        _build_rebar_last_point_annotation,
        _pick_extreme_rebar_layers,
    )

    catalog = load_material_catalog()
    section = build_section_input_from_draft(default_draft_inputs(), catalog)
    result = solve_bending_capacity(section, catalog)
    last_point = result.curve_points[-1]
    top_rebar, bottom_rebar = _pick_extreme_rebar_layers(section)

    concrete_annotation = _build_concrete_last_point_annotation(result)
    top_annotation = _build_rebar_last_point_annotation(section, result, rebar_index=top_rebar[0])
    bottom_annotation = _build_rebar_last_point_annotation(section, result, rebar_index=bottom_rebar[0])

    top_expected = (
        last_point.top_strain
        + (last_point.bottom_strain - last_point.top_strain) * top_rebar[1].z_mm / section.section_height_mm
    ) * 100000.0
    bottom_expected_raw = (
        last_point.top_strain
        + (last_point.bottom_strain - last_point.top_strain) * bottom_rebar[1].z_mm / section.section_height_mm
    ) * 100000.0
    bottom_expected = -bottom_expected_raw

    assert concrete_annotation.label == "εmax"
    assert concrete_annotation.target_strain == pytest.approx(last_point.top_strain * 100000.0)
    assert concrete_annotation.moment_kNm == pytest.approx(last_point.moment_kNm)
    assert concrete_annotation.within_chart_range is True
    assert top_annotation.label == "εmax"
    assert top_annotation.target_strain == pytest.approx(top_expected)
    assert top_annotation.moment_kNm == pytest.approx(last_point.moment_kNm)
    assert top_annotation.within_chart_range is True
    assert bottom_annotation.label == "εmax"
    assert bottom_annotation.target_strain == pytest.approx(bottom_expected)
    assert bottom_annotation.moment_kNm == pytest.approx(last_point.moment_kNm)
    assert bottom_annotation.within_chart_range is True


def test_moment_strain_annotation_layers_include_both_guides_for_in_range_point():
    from streamlit_app import _build_moment_strain_annotation_layers

    annotation = ChartLimitAnnotation(
        label="ε_yk",
        target_strain=-281.0,
        moment_kNm=11.03,
        within_chart_range=True,
    )

    layers = _build_moment_strain_annotation_layers(
        annotation,
        x_field="ε_s, 10^-5",
        x_axis_origin=-1272.2,
        x_domain_min=-1272.2,
        x_domain_max=0.0,
        y_axis_origin=0.0,
        y_label_value=11.07,
        color="#8c5a3a",
        label_text="ε_yk = -281.00 ·10^-5; M = 11.03 кН·м",
    )

    layer_specs = [layer.to_dict() for layer in layers]
    text_dataset = next(iter(layer_specs[3]["datasets"].values()))

    assert len(layer_specs) == 4
    assert [spec["mark"]["type"] for spec in layer_specs] == ["rule", "rule", "point", "text"]
    assert "x2" in layer_specs[0]["encoding"]
    assert "y2" in layer_specs[0]["encoding"]
    assert "x2" in layer_specs[1]["encoding"]
    assert "y2" in layer_specs[1]["encoding"]
    assert text_dataset[0]["Підпис"] == "ε_yk = -281.00 ·10^-5; M = 11.03 кН·м"


def test_moment_strain_annotation_layers_shift_right_edge_label_inside_chart():
    from streamlit_app import _build_moment_strain_annotation_layers

    annotation = ChartLimitAnnotation(
        label="εmax",
        target_strain=257.0,
        moment_kNm=11.01,
        within_chart_range=True,
    )

    layers = _build_moment_strain_annotation_layers(
        annotation,
        x_field="ε_c,top, 10^-5",
        x_axis_origin=0.0,
        x_domain_min=0.0,
        x_domain_max=257.0,
        y_axis_origin=0.0,
        y_label_value=11.07,
        color="#334155",
        label_text="εmax = 257.00 ·10^-5; M = 11.01 кН·м",
    )

    text_mark = layers[3].to_dict()["mark"]

    assert text_mark["align"] == "right"
    assert text_mark["dx"] < 0


def test_moment_strain_annotation_layers_can_place_max_label_below_point():
    from streamlit_app import _build_moment_strain_annotation_layers

    annotation = ChartLimitAnnotation(
        label="εmax",
        target_strain=257.0,
        moment_kNm=11.01,
        within_chart_range=True,
    )

    layers = _build_moment_strain_annotation_layers(
        annotation,
        x_field="ε_c,top, 10^-5",
        x_axis_origin=0.0,
        x_domain_min=0.0,
        x_domain_max=257.0,
        y_axis_origin=0.0,
        y_label_value=11.07,
        color="#334155",
        label_position="below",
        label_text="εmax = 257.00 ·10^-5; M = 11.01 кН·м",
    )

    text_mark = layers[3].to_dict()["mark"]

    assert text_mark["baseline"] == "top"
    assert text_mark["dy"] > 0


def test_moment_strain_annotation_layers_keep_only_vertical_boundary_for_out_of_range_limit():
    from streamlit_app import _build_moment_strain_annotation_layers

    annotation = ChartLimitAnnotation(
        label="ε_cu1,ck",
        target_strain=263.0,
        moment_kNm=None,
        within_chart_range=False,
    )

    layers = _build_moment_strain_annotation_layers(
        annotation,
        x_field="ε_c,top, 10^-5",
        x_axis_origin=0.0,
        x_domain_min=0.0,
        x_domain_max=257.0,
        y_axis_origin=0.0,
        y_label_value=11.07,
        color="#8c5a3a",
        label_text="ε_cu1,ck = 263.00 ·10^-5",
    )

    layer_specs = [layer.to_dict() for layer in layers]
    text_dataset = next(iter(layer_specs[1]["datasets"].values()))

    assert len(layer_specs) == 2
    assert [spec["mark"]["type"] for spec in layer_specs] == ["rule", "text"]
    assert text_dataset[0]["Підпис"] == "ε_cu1,ck = 263.00 ·10^-5"
    assert layer_specs[1]["mark"]["align"] == "right"
    assert layer_specs[1]["mark"]["dx"] < 0


def test_moment_strain_annotation_layers_hide_text_when_overlay_mode_is_none():
    from streamlit_app import _build_moment_strain_annotation_layers

    annotation = ChartLimitAnnotation(
        label="εmax",
        target_strain=-354.68,
        moment_kNm=11.01,
        within_chart_range=True,
    )

    layers = _build_moment_strain_annotation_layers(
        annotation,
        x_field="ε_s, 10^-5",
        x_axis_origin=-1272.2,
        x_domain_min=-1272.2,
        x_domain_max=0.0,
        y_axis_origin=0.0,
        y_label_value=11.07,
        color="#334155",
        text_mode="none",
    )

    layer_specs = [layer.to_dict() for layer in layers]

    assert len(layer_specs) == 3
    assert [spec["mark"]["type"] for spec in layer_specs] == ["rule", "rule", "point"]


def test_concrete_chart_explanation_html_describes_normative_limit_and_reached_point():
    from streamlit_app import _build_concrete_chart_explanation_html

    limit_annotation = ChartLimitAnnotation(
        label="ε_cu1,ck",
        target_strain=263.0,
        moment_kNm=None,
        within_chart_range=False,
    )
    max_annotation = ChartLimitAnnotation(
        label="εmax",
        target_strain=257.0,
        moment_kNm=11.01,
        within_chart_range=True,
    )

    html = _build_concrete_chart_explanation_html(
        chart_title="Момент-деформація бетону",
        concrete_class="C40/50",
        limit_annotation=limit_annotation,
        max_annotation=max_annotation,
    )

    assert 'data-role="chart-explanation-card"' in html
    assert 'data-role="chart-explanation-title"' in html
    assert "Момент-деформація бетону" in html
    assert "C40/50" in html
    assert "ε_cu1,ck = 263.00 ·10^-5" in html
    assert "нормативна межа" in html
    assert "не потрапила в побудовану криву" in html
    assert "εmax = 257.00 ·10^-5" in html
    assert "M(εmax) = 11.01 кН·м" in html
    assert "остання фактично досягнута точка" in html
    assert "solver" not in html


def test_rebar_chart_explanation_html_describes_limit_and_epsilon_max():
    from streamlit_app import _build_rebar_chart_explanation_html

    limit_annotation = ChartLimitAnnotation(
        label="ε_yk",
        target_strain=-281.0,
        moment_kNm=11.04,
        within_chart_range=True,
    )
    max_annotation = ChartLimitAnnotation(
        label="εmax",
        target_strain=-354.68,
        moment_kNm=11.01,
        within_chart_range=True,
    )

    html = _build_rebar_chart_explanation_html(
        chart_title="Момент-деформація верхньої арматури",
        rebar_label="A1",
        steel_class="A500C",
        z_mm=40.0,
        limit_annotation=limit_annotation,
        max_annotation=max_annotation,
    )

    assert 'data-role="chart-explanation-card"' in html
    assert "Момент-деформація верхньої арматури" in html
    assert "A1" in html
    assert "A500C" in html
    assert "z = 40.0 мм" in html
    assert "ε_yk" in html
    assert "деформація текучості" in html
    assert "M(ε_yk) = 11.04 кН·м" in html
    assert "εmax = -354.68 ·10^-5" in html
    assert "M(εmax) = 11.01 кН·м" in html
    assert "від’ємній гілці" in html
    assert "solver" not in html


def test_moment_strain_chart_keeps_x_axis_left_to_right_for_decreasing_rebar_strains():
    from streamlit_app import _build_moment_strain_chart

    chart_df = pd.DataFrame(
        {
            "Крок": [3, 2, 1],
            "M, кН·м": [3.0, 2.0, 0.0],
            "ε_s, 10^-5": [-30.0, -10.0, 0.0],
        }
    )

    chart = _build_moment_strain_chart(
        chart_df,
        x_field="ε_s, 10^-5",
        selected_step=3,
        tooltip_fields=["Крок", "ε_s, 10^-5", "M, кН·м"],
        annotations=[],
    )

    base_encoding = chart.to_dict()["layer"][0]["encoding"]

    assert base_encoding["x"]["scale"]["reverse"] is False
    assert base_encoding["order"]["field"] == "Крок"


def test_moment_strain_chart_reverses_top_rebar_axis_for_engineering_presentation():
    from streamlit_app import _build_moment_strain_chart

    chart_df = pd.DataFrame(
        {
            "Крок": [3, 2, 1],
            "M, кН·м": [3.0, 2.0, 0.0],
            "ε_s, 10^-5": [-30.0, -10.0, 0.0],
        }
    )

    chart = _build_moment_strain_chart(
        chart_df,
        x_field="ε_s, 10^-5",
        selected_step=3,
        tooltip_fields=["Крок", "ε_s, 10^-5", "M, кН·м"],
        annotations=[],
        chart_id="top_rebar_moment_strain",
    )

    base_encoding = chart.to_dict()["layer"][0]["encoding"]

    assert base_encoding["x"]["scale"]["reverse"] is True
    assert base_encoding["order"]["field"] == "Крок"


def test_moment_strain_chart_adds_experimental_overlay_layer():
    from streamlit_app import _build_moment_strain_chart

    chart_df = pd.DataFrame(
        {
            "Крок": [1, 2, 3],
            "M, кН·м": [0.0, 30.0, 60.0],
            "ε_c,top, 10^-5": [0.0, 100.0, 200.0],
        }
    )
    experimental_df = pd.DataFrame(
        {
            "ε_c,top, 10^-5": [0.0, 90.0, 180.0],
            "M, кН·м": [0.0, 28.0, 58.0],
        }
    )

    chart = _build_moment_strain_chart(
        chart_df,
        x_field="ε_c,top, 10^-5",
        selected_step=3,
        tooltip_fields=["Крок", "ε_c,top, 10^-5", "M, кН·м"],
        annotations=[],
        experimental_df=experimental_df,
    )
    layers = chart.to_dict()["layer"]
    experimental_layers = [layer for layer in layers if layer.get("mark", {}).get("strokeDash") == [8, 4]]

    assert len(experimental_layers) == 1
    assert experimental_layers[0]["mark"]["color"] == "#b45309"


def test_build_theoretical_state_at_deflection_interpolates_moment_and_strains_in_step_order():
    from streamlit_app import _build_theoretical_state_at_deflection

    theory_df = pd.DataFrame(
        {
            "Крок": [1, 2, 3],
            "M, кН·м": [0.0, 10.0, 20.0],
            "f, мм": [0.0, 5.0, 10.0],
            "ε_c,top, 10^-5": [0.0, 100.0, 200.0],
            "ε_s,top, 10^-5": [0.0, -150.0, -300.0],
            "ε_s,bot, 10^-5": [0.0, 180.0, 360.0],
        }
    )

    reference = _build_theoretical_state_at_deflection(theory_df, target_deflection_mm=7.5)

    assert reference.deflection_mm == pytest.approx(7.5)
    assert reference.moment_kNm == pytest.approx(15.0)
    assert reference.concrete_strain_e5 == pytest.approx(150.0)
    assert reference.top_rebar_strain_e5 == pytest.approx(-225.0)
    assert reference.bottom_rebar_strain_e5 == pytest.approx(270.0)
    assert reference.within_range is True


def test_build_experimental_deflection_annotation_reports_out_of_range_state():
    from streamlit_app import _build_experimental_deflection_annotation

    experimental_df = pd.DataFrame(
        {
            "f, мм": [0.0, 3.0, 6.0],
            "M, кН·м": [0.0, 12.0, 20.0],
        }
    )

    annotation = _build_experimental_deflection_annotation(experimental_df, target_deflection_mm=7.5)

    assert annotation.label == "exp @ f_u"
    assert annotation.target_strain == pytest.approx(7.5)
    assert annotation.moment_kNm is None
    assert annotation.within_chart_range is False


def test_build_experimental_primary_chart_supports_indic_and_dic_overlays_and_characteristic_labels():
    from streamlit_app import ChartAnnotationOverlay, _build_experimental_primary_chart

    indic_df = pd.DataFrame(
        {
            "f, мм": [0.0, 4.0, 8.0],
            "M, кН·м": [0.0, 11.0, 19.0],
        }
    )
    dic_df = pd.DataFrame(
        {
            "f, мм": [0.0, 4.5, 8.5],
            "M, кН·м": [0.0, 12.5, 18.0],
        }
    )
    theory_df = pd.DataFrame(
        {
            "Крок": [1, 2, 3],
            "f, мм": [0.0, 5.0, 10.0],
            "M, кН·м": [0.0, 10.0, 20.0],
        }
    )

    chart = _build_experimental_primary_chart(
        x_field="f, мм",
        indic_df=indic_df,
        dic_df=dic_df,
        theory_df=theory_df,
        show_theory=True,
        show_indic=True,
        show_dic=True,
        annotations=[
            ChartAnnotationOverlay(
                ChartLimitAnnotation(
                    label="theory @ f_u",
                    target_strain=7.5,
                    moment_kNm=15.0,
                    within_chart_range=True,
                ),
                "#1d4ed8",
                label_text="Theory @ f_u",
            ),
            ChartAnnotationOverlay(
                ChartLimitAnnotation(
                    label="indic @ f_u",
                    target_strain=7.5,
                    moment_kNm=18.0,
                    within_chart_range=True,
                ),
                "#b45309",
                label_position="below",
                label_text="Indic @ f_u",
            ),
            ChartAnnotationOverlay(
                ChartLimitAnnotation(
                    label="dic @ f_u",
                    target_strain=7.5,
                    moment_kNm=17.0,
                    within_chart_range=True,
                ),
                "#059669",
                label_position="below",
                label_text="DIC @ f_u",
            ),
        ],
    )

    layer_specs = chart.to_dict()["layer"]
    datasets = chart.to_dict().get("datasets", {})
    line_layers = [layer for layer in layer_specs if layer.get("mark", {}).get("type") == "line"]
    point_layers = [layer for layer in layer_specs if layer.get("mark", {}).get("type") == "point"]
    text_layers = [layer for layer in layer_specs if layer.get("mark", {}).get("type") == "text"]

    assert any(layer["mark"]["color"] == "#b45309" and "strokeDash" not in layer["mark"] for layer in line_layers)
    assert any(layer["mark"]["color"] == "#059669" and "strokeDash" not in layer["mark"] for layer in line_layers)
    assert any(layer["mark"]["color"] == "#1d4ed8" and layer["mark"]["strokeDash"] == [6, 4] for layer in line_layers)
    assert sum(1 for layer in point_layers if layer["mark"]["color"] in {"#1d4ed8", "#b45309", "#059669"}) >= 3
    assert any(
        row["Підпис"] == "Theory @ f_u"
        for layer in text_layers
        for row in datasets.get(layer.get("data", {}).get("name", ""), [])
    )
    assert any(
        row["Підпис"] == "Indic @ f_u"
        for layer in text_layers
        for row in datasets.get(layer.get("data", {}).get("name", ""), [])
    )
    assert any(
        row["Підпис"] == "DIC @ f_u"
        for layer in text_layers
        for row in datasets.get(layer.get("data", {}).get("name", ""), [])
    )


def test_build_experimental_primary_chart_reverses_top_rebar_axis_for_theory_and_overlays():
    from streamlit_app import _build_experimental_primary_chart

    indic_df = pd.DataFrame(
        {
            "ε_s, 10^-5": [-420.0, -120.0, 0.0],
            "M, кН·м": [30.0, 10.0, 0.0],
        }
    )
    dic_df = pd.DataFrame(
        {
            "ε_s, 10^-5": [-430.0, -130.0, 0.0],
            "M, кН·м": [28.0, 9.0, 0.0],
        }
    )
    theory_df = pd.DataFrame(
        {
            "Крок": [3, 2, 1],
            "ε_s, 10^-5": [-360.0, -180.0, 0.0],
            "M, кН·м": [11.0, 6.0, 0.0],
        }
    )

    chart = _build_experimental_primary_chart(
        chart_id="top_rebar_moment_strain",
        x_field="ε_s, 10^-5",
        indic_df=indic_df,
        dic_df=dic_df,
        theory_df=theory_df,
        show_theory=True,
        show_indic=True,
        show_dic=True,
    )

    x_encodings = [
        layer["encoding"]["x"]["scale"]["reverse"]
        for layer in chart.to_dict()["layer"]
        if layer.get("mark", {}).get("type") == "line"
    ]

    assert x_encodings
    assert all(value is True for value in x_encodings)


def test_build_experimental_reference_summary_html_renders_deflection_and_material_values():
    from streamlit_app import (
        _build_experimental_reference_summary_html,
        _build_theoretical_state_at_deflection,
    )

    theory_df = pd.DataFrame(
        {
            "Крок": [1, 2, 3],
            "M, кН·м": [0.0, 10.0, 20.0],
            "f, мм": [0.0, 5.0, 10.0],
            "ε_c,top, 10^-5": [0.0, 100.0, 200.0],
            "ε_s,top, 10^-5": [0.0, -150.0, -300.0],
            "ε_s,bot, 10^-5": [0.0, 180.0, 360.0],
        }
    )
    theoretical_state = _build_theoretical_state_at_deflection(theory_df, target_deflection_mm=7.5)
    indic_annotation = ChartLimitAnnotation(
        label="indic @ f_u",
        target_strain=7.5,
        moment_kNm=18.0,
        within_chart_range=True,
    )
    dic_annotation = ChartLimitAnnotation(
        label="dic @ f_u",
        target_strain=7.5,
        moment_kNm=16.5,
        within_chart_range=True,
    )
    limit_annotation = ChartLimitAnnotation(
        label="ε_yk",
        target_strain=-281.0,
        moment_kNm=11.04,
        within_chart_range=True,
    )
    max_annotation = ChartLimitAnnotation(
        label="εmax",
        target_strain=-354.68,
        moment_kNm=11.01,
        within_chart_range=True,
    )

    html = _build_experimental_reference_summary_html(
        chart_id="deflection_mf",
        target_deflection_mm=7.5,
        theoretical_state=theoretical_state,
        indic_annotation=indic_annotation,
        dic_annotation=dic_annotation,
        limit_annotation=limit_annotation,
        max_annotation=max_annotation,
    )

    assert 'data-role="experimental-reference-card"' in html
    assert "M_theory(f_u) = 15.00 кН·м" in html
    assert "M_Indic(f_u) = 18.00 кН·м" in html
    assert "M_DIC(f_u) = 16.50 кН·м" in html
    assert "f_u = 7.50 мм" in html

    concrete_html = _build_experimental_reference_summary_html(
        chart_id="concrete_moment_strain",
        target_deflection_mm=7.5,
        theoretical_state=theoretical_state,
        indic_annotation=None,
        dic_annotation=None,
        limit_annotation=ChartLimitAnnotation(
            label="ε_cu1,ck",
            target_strain=263.0,
            moment_kNm=11.06,
            within_chart_range=True,
        ),
        max_annotation=ChartLimitAnnotation(
            label="εmax",
            target_strain=257.0,
            moment_kNm=11.01,
            within_chart_range=True,
        ),
    )

    assert "ε_c(theory @ f_u) = 150.00 ·10^-5" in concrete_html
    assert "M_theory(f_u) = 15.00 кН·м" in concrete_html
    assert "M(ε_cu1,ck) = 11.06 кН·м" in concrete_html


def test_build_experimental_fu_summary_table_html_renders_columns_and_out_of_range_fallback():
    from streamlit_app import _build_experimental_fu_summary_table_html

    theoretical_state = type(
        "TheoryState",
        (),
        {"moment_kNm": 15.0},
    )()
    indic_annotation = ChartLimitAnnotation(
        label="indic @ f_u",
        target_strain=7.5,
        moment_kNm=18.0,
        within_chart_range=True,
    )
    dic_annotation = ChartLimitAnnotation(
        label="dic @ f_u",
        target_strain=7.5,
        moment_kNm=None,
        within_chart_range=False,
    )

    html = _build_experimental_fu_summary_table_html(
        target_deflection_mm=7.5,
        theoretical_state=theoretical_state,
        indic_annotation=indic_annotation,
        dic_annotation=dic_annotation,
        include_dic=True,
    )

    assert 'data-role="experimental-fu-summary"' in html
    assert "Граничний момент при досягненні граничного прогину" in html
    assert "<table" in html
    assert "f_u, мм" in html
    assert "M_theory(f_u), кН·м" in html
    assert "M_Indic(f_u), кН·м" in html
    assert "M_DIC(f_u), кН·м" in html
    assert "7.50" in html
    assert "15.00" in html
    assert "18.00" in html
    assert "поза діапазоном експерименту" in html


def test_build_experimental_explanation_html_uses_article_notation_for_indic_and_dic():
    from streamlit_app import _build_experimental_explanation_html

    html = _build_experimental_explanation_html(
        chart_id="deflection_mf",
        reference_mode=False,
        has_indic=True,
        has_dic=True,
    )

    assert 'data-role="experimental-explanation-card"' in html
    assert "Пояснення позначень" in html
    assert "Defl_Indic" in html
    assert "Defl_DIC" in html
    assert "M_ULS" in html
    assert "M_max" in html
    assert "f_u" in html


def test_analytics_summary_table_uses_curve_and_extreme_rebar_layers():
    from streamlit_app import _build_analytics_summary_df

    catalog = load_material_catalog()
    section = make_case()
    result = solve_bending_capacity(section, catalog)

    summary_df = _build_analytics_summary_df(section, result)
    first_point = result.curve_points[0]
    top_expected = first_point.top_strain + (
        (first_point.bottom_strain - first_point.top_strain) * section.rebar_layers[0].z_mm / section.section_height_mm
    )
    bottom_expected_raw = first_point.top_strain + (
        (first_point.bottom_strain - first_point.top_strain) * section.rebar_layers[1].z_mm / section.section_height_mm
    )
    bottom_expected = -bottom_expected_raw

    assert list(summary_df.columns) == [
        "Крок",
        "M, кН·м",
        "κ, 1/м",
        "ε_c,top, 10^-5",
        "ε_s,top, 10^-5",
        "ε_s,bot, 10^-5",
    ]
    assert len(summary_df) == len(result.curve_points)
    assert summary_df.iloc[0]["Крок"] == first_point.step_index
    assert summary_df.iloc[0]["M, кН·м"] == pytest.approx(first_point.moment_kNm)
    assert summary_df.iloc[0]["κ, 1/м"] == pytest.approx(first_point.curvature_1_per_m)
    assert summary_df.iloc[0]["ε_c,top, 10^-5"] == pytest.approx(first_point.top_strain * 100000.0)
    assert summary_df.iloc[0]["ε_s,top, 10^-5"] == pytest.approx(top_expected * 100000.0)
    assert summary_df.iloc[0]["ε_s,bot, 10^-5"] == pytest.approx(bottom_expected * 100000.0)


def test_reference_single_rebar_chart_uses_positive_bottom_display_branch():
    from streamlit_app import _build_rebar_last_point_annotation, _build_rebar_moment_strain_df

    catalog = load_material_catalog()
    draft = default_draft_inputs()
    draft["has_strengthening_layer"] = False
    draft["rebar_layers"] = [draft["rebar_layers"][1]]
    section = build_section_input_from_draft(draft, catalog)
    result = solve_bending_capacity(section, catalog)

    rebar_df = _build_rebar_moment_strain_df(section, result, rebar_index=1)
    last_annotation = _build_rebar_last_point_annotation(section, result, rebar_index=1)

    assert rebar_df["ε_s, 10^-5"].iloc[0] == pytest.approx(0.0)
    assert rebar_df["ε_s, 10^-5"].is_monotonic_increasing
    assert rebar_df["ε_s, 10^-5"].iloc[-1] > 0.0
    assert last_annotation.target_strain > 0.0


def test_termination_summary_html_explains_stop_reason_and_preceding_state():
    from streamlit_app import _build_termination_summary_html

    catalog = load_material_catalog()
    result = solve_bending_capacity(make_case(), catalog, outer_steps=12)

    html = _build_termination_summary_html(result)

    assert 'data-role="termination-panel"' in html
    assert 'data-role="termination-value-reason"' in html
    assert 'data-role="termination-value-last-step"' in html
    assert 'data-role="termination-value-last-moment"' in html
    assert 'data-role="termination-value-previous-step"' in html
    assert 'data-role="termination-help-note"' in html
    assert 'data-role="termination-narrative"' in html
    assert str(result.termination.last_step) in html
    assert f"{result.termination.last_moment_kNm:.2f}" in html
    assert (
        f"{result.termination.attempted_top_strain * 100000.0:.2f}" in html
        if result.termination.attempted_top_strain is not None
        else True
    )
    assert (
        "Спроба продовжити криву" in html
        if result.termination.attempted_top_strain is not None
        else True
    )
    assert "solver" not in html
    assert "не застосовується" not in html


def test_termination_reason_descriptions_use_engineering_language():
    from streamlit_app import _describe_termination_reason

    assert "критерієм бетону" in _describe_termination_reason("completed_at_concrete_limit")
    assert "арматури" in _describe_termination_reason("no_tension_equilibrium_at_steel_limit")
    assert "стискувальна" in _describe_termination_reason("no_compression_equilibrium_at_upper_bound")
    residual_text = _describe_termination_reason("high_axial_residual_after_iteration")
    assert "числова" in residual_text
    assert "не матеріальне руйнування" in residual_text


def test_termination_narrative_distinguishes_numerical_stop_from_material_failure():
    from streamlit_app import _build_termination_narrative

    catalog = load_material_catalog()
    result = solve_bending_capacity(make_case(), catalog, outer_steps=12)
    adjusted_result = replace(
        result,
        termination=replace(
            result.termination,
            reason_code="high_axial_residual_after_iteration",
            attempted_step=result.termination.last_step + 1,
            residual_kN=6.25,
        ),
    )

    narrative = _build_termination_narrative(adjusted_result)

    assert "не матеріальне руйнування" in narrative
    assert "числова збіжність" in narrative
    assert str(adjusted_result.termination.attempted_step) in narrative


def test_annotation_layers_apply_manual_dx_and_dy_offsets():
    from streamlit_app import _build_moment_strain_annotation_layers

    annotation = ChartLimitAnnotation(label="ε_test", target_strain=120.0, moment_kNm=85.0, within_chart_range=True)

    base_layers = _build_moment_strain_annotation_layers(
        annotation,
        x_field="ε_c,top, 10^-5",
        x_axis_origin=0.0,
        x_domain_min=0.0,
        x_domain_max=300.0,
        y_axis_origin=0.0,
        y_label_value=100.0,
    )
    shifted_layers = _build_moment_strain_annotation_layers(
        annotation,
        x_field="ε_c,top, 10^-5",
        x_axis_origin=0.0,
        x_domain_min=0.0,
        x_domain_max=300.0,
        y_axis_origin=0.0,
        y_label_value=100.0,
        manual_dx=14,
        manual_dy=-9,
    )

    base_text_mark = base_layers[-1].to_dict()["mark"]
    shifted_text_mark = shifted_layers[-1].to_dict()["mark"]

    assert shifted_text_mark["dx"] == base_text_mark["dx"] + 14
    assert shifted_text_mark["dy"] == base_text_mark["dy"] - 9


def test_draft_validation_requires_outer_steps_in_supported_range():
    catalog = load_material_catalog()
    too_small = default_draft_inputs()
    too_small["outer_steps"] = 1

    too_large = default_draft_inputs()
    too_large["outer_steps"] = 201

    not_integer = default_draft_inputs()
    not_integer["outer_steps"] = 12.5

    too_small_errors = validate_draft_inputs(too_small, catalog)
    too_large_errors = validate_draft_inputs(too_large, catalog)
    not_integer_errors = validate_draft_inputs(not_integer, catalog)

    assert any("Кількість кроків" in error for error in too_small_errors)
    assert any("Кількість кроків" in error for error in too_large_errors)
    assert any("Кількість кроків" in error for error in not_integer_errors)


def test_streamlit_app_renders_without_exception():
    app_path = Path("streamlit_app.py")
    assert app_path.exists(), "streamlit_app.py must exist"

    app = AppTest.from_file(str(app_path))
    app.run(timeout=10)

    workspace_root = app.main

    assert len(app.exception) == 0
    assert _find_widget_by_label(app.button, "Переріз").label == "Переріз"
    assert _find_widget_by_label(app.button, "II ГГС").label == "II ГГС"
    assert _find_widget_by_label(app.button, "Експеримент").label == "Експеримент"
    hero = _find_markdown_containing(app, 'data-role="hero-banner"')
    assert "Розрахунок згину залізобетонного перерізу" in hero
    assert "ДСТУ / ДБН" in hero
    assert "Переріз" in hero
    assert "II ГГС" in hero
    assert "Експеримент" in hero
    toolbar = _find_markdown_containing(app, '<section class="workspace-nav-toolbar" data-role="workspace-nav-toolbar">')
    assert 'data-role="workspace-toolbar-chip-grid"' in toolbar
    assert 'data-role="workspace-toolbar-chip"' in toolbar
    assert 'data-role="workspace-toolbar-note"' in toolbar
    assert not _has_markdown_containing(app, 'data-role="workspace-launcher"')
    assert not _has_markdown_containing(app, 'data-role="shared-state-panel"')
    assert not any(
        getattr(button, "label", None) in {"Відкрити Переріз", "Відкрити II ГГС", "Відкрити Експеримент"}
        for button in app.button
    )
    assert not any("workspace_view" in getattr(warning, "value", "") for warning in getattr(app, "warning", []))
    assert _find_widget_by_label(app.slider, "Активна точка кривої").label == "Активна точка кривої"
    assert _find_widget_by_label(app.number_input, "Висота перерізу h, мм").value == 120.0
    assert _find_widget_by_label(app.number_input, "Ширина перерізу b, мм").value == 500.0
    assert _find_widget_by_label(app.number_input, "Кількість кроків розрахунку").value == 12
    assert _find_widget_by_label(app.button, "Перерахувати").label == "Перерахувати"
    assert _find_metric(app, "Несуча здатність M_Rd, кН·м").label == "Несуча здатність M_Rd, кН·м"
    assert _find_metric(app, "Кривизна κ_peak, 1/м").label == "Кривизна κ_peak, 1/м"
    assert not _has_metric(app, "Нев'язка ΣN, кН")
    assert not any(getattr(button, "label", None) == "Додати шар арматури" for button in app.button)
    assert not any(getattr(button, "label", None) == "Видалити" for button in app.button)
    assert any(getattr(subheader, "value", None) == "Діаграма M-κ" for subheader in app.subheader)
    assert any(getattr(subheader, "value", None) == "Момент-деформація бетону" for subheader in app.subheader)
    assert any(getattr(subheader, "value", None) == "Момент-деформація верхньої арматури" for subheader in app.subheader)
    assert any(getattr(subheader, "value", None) == "Момент-деформація нижньої арматури" for subheader in app.subheader)
    assert not any(getattr(subheader, "value", None) == "Графік деформацій бетону" for subheader in app.subheader)
    assert not any(getattr(subheader, "value", None) == "Проміжні результати" for subheader in app.subheader)
    assert not any(getattr(subheader, "value", None) == "Зусилля в шарах для обраної точки" for subheader in app.subheader)
    point_panel = _find_markdown_containing(app, 'data-role="point-chip-grid"')
    force_panel = _find_markdown_containing(app, 'data-role="force-card-grid"')
    termination_panel = _find_markdown_containing(app, 'data-role="termination-panel"')
    assert 'data-role="point-value-step"' in point_panel
    assert 'data-role="point-value-moment"' in point_panel
    assert 'data-role="point-value-curvature"' in point_panel
    assert 'data-role="point-value-neutral-axis"' in point_panel
    assert 'data-role="point-value-residual"' in point_panel
    assert 'data-role="force-card"' in force_panel
    assert 'data-role="termination-value-reason"' in termination_panel
    assert 'data-role="termination-help-note"' in termination_panel
    assert 'data-role="termination-narrative"' in termination_panel
    assert not _has_markdown_containing(app, 'data-role="chart-limit-badges"')
    assert _count_markdown_containing(app, 'data-role="chart-explanation-card"') >= 3
    assert _has_markdown_containing(app, "нормативна гранична деформація стиснутого бетону")
    assert _has_markdown_containing(app, "деформація текучості")
    assert _has_markdown_containing(app, "від’ємній гілці")
    assert _has_markdown_containing(app, "κ показує кривизну перерізу")
    assert _has_markdown_containing(app, "εmax")
    assert _has_markdown_containing(app, "M(εmax)")
    assert _has_markdown_containing(app, "ε_cu1,ck = 263.00 ·10^-5")
    assert not _has_markdown_containing(app, "solver")
    assert not _has_markdown_containing(app, "не застосовується")
    concrete_preview = _find_dataframe(app, ["Шар", "b, мм", "h, мм", "Клас", "Статус"])
    rebar_preview = _find_dataframe(app, ["Шар", "Грань", "a, мм", "z, мм", "n, шт.", "d, мм", "Клас", "Статус"])
    assert concrete_preview.iloc[0]["h, мм"] == 60.0
    assert concrete_preview.iloc[1]["Клас"] == "C40/50"
    assert rebar_preview.iloc[0]["Клас"] == "A500C"
    assert rebar_preview.iloc[0]["n, шт."] == 4
    assert rebar_preview.iloc[0]["d, мм"] == 8
    assert rebar_preview.iloc[1]["z, мм"] == 100.0
    summary_table = _find_dataframe(app, [
        "Крок",
        "M, кН·м",
        "κ, 1/м",
        "ε_c,top, 10^-5",
        "ε_s,top, 10^-5",
        "ε_s,bot, 10^-5",
    ])
    assert list(summary_table.columns) == [
        "Крок",
        "M, кН·м",
        "κ, 1/м",
        "ε_c,top, 10^-5",
        "ε_s,top, 10^-5",
        "ε_s,bot, 10^-5",
    ]
    assert summary_table.iloc[0]["Крок"] == 1
    assert summary_table.iloc[0]["M, кН·м"] == pytest.approx(0.0)
    assert len(summary_table) <= 12
    assert not _has_dataframe(app, ["Крок", "M, кН·м", "κ, 1/м", "ε_c,top, ‰", "ε_c,bot, ‰"])
    assert not _has_dataframe(app, ["Крок", "M, кН·м", "ε_c,top, 10^-5"])
    assert not _has_dataframe(app, ["z, мм", "ε_c, ‰"])
    assert not _has_dataframe(app, ["Тип", "№", "Клас", "z, мм", "A, мм²", "ε, ‰", "σ, МПа", "N, кН"])
    drawing = _find_markdown_containing(app, "<svg")
    assert "Креслення перерізу" in drawing
    assert "2-га форма рівноваги" in drawing
    assert "Активна точка" in drawing
    assert "1-ша форма рівноваги" not in drawing
    assert "форма не реалізована для цього набору даних" not in drawing
    assert "A1: 4 x 8" in drawing
    assert "A500C" in drawing
    assert "N_A1 =" in drawing
    assert 'data-role="metric-legend"' in drawing
    assert 'data-role="metric-chip"' in drawing
    assert 'data-role="form-grid-line"' in drawing
    assert 'data-role="top-row-layout"' in drawing
    assert 'data-role="dimension-lane-left-h"' in drawing
    assert 'data-role="dimension-lane-left-z1"' in drawing
    assert 'data-role="dimension-lane-left-z2"' in drawing
    assert 'data-role="dimension-lane-right-h1"' in drawing
    assert 'data-role="dimension-lane-right-h2"' in drawing
    assert 'data-role="dimension-label-chip"' in drawing
    assert _has_markdown_containing_in_node(workspace_root, 'data-role="results-section-lead"')
    assert _has_markdown_containing_in_node(workspace_root, 'data-role="drawing-showcase"')
    assert _has_markdown_containing_in_node(workspace_root, 'data-role="termination-panel"')
    assert _has_markdown_containing_in_node(workspace_root, 'data-role="input-section-lead"')
    assert _has_markdown_containing_in_node(
        workspace_root,
        'data-role="workspace-full-width-layout" data-workspace="section"',
    )
    assert _has_markdown_containing_in_node(
        workspace_root,
        'data-role="workspace-input-panel" data-workspace="section"',
    )
    assert _has_markdown_containing_in_node(
        workspace_root,
        'data-role="workspace-results-panel" data-workspace="section"',
    )
    assert not _has_markdown_containing(app, 'data-role="serviceability-section-lead"')
    assert not _has_markdown_containing(app, 'data-role="experimental-upload-status"')
    assert len(app.button_group) == 0


def test_streamlit_app_renders_serviceability_controls_and_theory():
    app = AppTest.from_file("streamlit_app.py")

    app.run(timeout=10)

    serviceability_tab = _activate_workspace(app, "serviceability")

    assert len(app.exception) == 0
    assert _find_widget_by_label_in_node(serviceability_tab, "number_input", "Розрахунковий проліт l, мм").value == 3000.0
    assert (
        _find_widget_by_label_in_node(serviceability_tab, "selectbox", "Розрахункова схема для прогину").value
        == "Балка на двох опорах з рівномірно розподіленим навантаженням"
    )
    assert (
        _find_widget_by_label_in_node(serviceability_tab, "selectbox", "Нормативний профіль обмеження прогину").value
        == "Естетико-психологічні"
    )
    assert _find_widget_by_label_in_node(serviceability_tab, "button", "Оновити перевірку").label == "Оновити перевірку"
    assert any(getattr(subheader, "value", None) == "Перевірка тріщин та прогинів" for subheader in serviceability_tab.subheader)
    assert any(getattr(subheader, "value", None) == "Діаграма M-f" for subheader in serviceability_tab.subheader)
    assert any(getattr(expander, "label", None) == "Пояснення та нормативна база" for expander in serviceability_tab.expander)
    assert _has_markdown_containing_in_node(
        serviceability_tab,
        'data-role="workspace-full-width-layout" data-workspace="serviceability"',
    )
    assert _has_markdown_containing_in_node(
        serviceability_tab,
        'data-role="workspace-input-panel" data-workspace="serviceability"',
    )
    assert _has_markdown_containing_in_node(
        serviceability_tab,
        'data-role="workspace-results-panel" data-workspace="serviceability"',
    )
    assert _has_markdown_containing_in_node(serviceability_tab, "ДСТУ Б В.2.6-156:2010")
    assert _has_markdown_containing_in_node(serviceability_tab, "ДСТУ Б В.1.2-3:2006")
    assert _has_markdown_containing_in_node(serviceability_tab, "вертикальні граничні прогини")
    assert _has_markdown_containing_in_node(serviceability_tab, "горизонтальні переміщення")
    assert _has_markdown_containing_in_node(serviceability_tab, "ухил покрівлі не менш як 1/200")
    assert not _has_markdown_containing_in_node(serviceability_tab, "(5.8)")
    assert not _has_markdown_containing_in_node(serviceability_tab, "(5.19)")
    assert not _has_markdown_containing_in_node(serviceability_tab, "w_k = s_r,max")
    assert not _has_markdown_containing_in_node(serviceability_tab, 'data-role="experimental-upload-status"')


def test_streamlit_app_places_workspace_toolbar_before_workspace_content():
    app = AppTest.from_file("streamlit_app.py")

    app.run(timeout=10)

    hero_index = _top_level_block_index_with_markdown(app, 'data-role="hero-banner"')
    toolbar_index = _top_level_block_index_with_markdown(app, '<section class="workspace-nav-toolbar" data-role="workspace-nav-toolbar">')
    toolbar_block = app.main.children[toolbar_index]
    input_block_index, input_child_index = _block_and_child_index_with_markdown_in_node(
        app.main,
        'data-role="input-section-lead"',
    )
    results_block_index, results_child_index = _block_and_child_index_with_markdown_in_node(
        app.main,
        'data-role="results-section-lead"',
    )

    assert len(app.exception) == 0
    assert hero_index < toolbar_index < input_block_index
    assert (input_block_index, input_child_index) < (results_block_index, results_child_index)
    assert _find_widget_by_label_in_node(toolbar_block, "button", "Переріз").label == "Переріз"
    assert _find_widget_by_label_in_node(toolbar_block, "button", "II ГГС").label == "II ГГС"
    assert _find_widget_by_label_in_node(toolbar_block, "button", "Експеримент").label == "Експеримент"
    assert _find_widget_by_label_in_node(toolbar_block, "slider", "Активна точка кривої").value >= 1


def test_streamlit_app_switches_workspace_view_and_preserves_active_point():
    app = AppTest.from_file("streamlit_app.py")

    app.run(timeout=10)
    _find_widget_by_label(app.slider, "Активна точка кривої").set_value(3)
    app.run(timeout=10)

    serviceability_root = _activate_workspace(app, "serviceability")

    assert len(app.exception) == 0
    assert _find_widget_by_label(app.slider, "Активна точка кривої").value == 3
    assert _has_markdown_containing_in_node(serviceability_root, 'data-role="serviceability-section-lead"')
    assert _has_markdown_containing_in_node(serviceability_root, 'data-role="serviceability-scheme-showcase"')
    assert not _has_markdown_containing_in_node(serviceability_root, 'data-role="results-section-lead"')

    experimental_root = _activate_workspace(app, "experimental")

    assert _find_widget_by_label(app.slider, "Активна точка кривої").value == 3
    assert _has_markdown_containing_in_node(experimental_root, 'data-role="experimental-section-lead"')
    assert _has_markdown_containing_in_node(experimental_root, 'data-role="experimental-upload-status"')
    assert not _has_markdown_containing_in_node(
        experimental_root,
        'data-role="workspace-input-panel" data-workspace="experimental"',
    )
    assert not _has_markdown_containing_in_node(
        experimental_root,
        'data-role="workspace-full-width-layout" data-workspace="experimental"',
    )


def test_deflection_curve_df_uses_all_curve_points_and_expected_columns():
    from streamlit_app import _build_deflection_curve_df

    catalog = load_material_catalog()
    draft = default_draft_inputs()
    section = build_section_input_from_draft(draft, catalog)
    result = solve_bending_capacity(section, catalog, outer_steps=int(draft["outer_steps"]))
    service_input = ServiceabilityInput(
        span_mm=6000.0,
        support_scheme="simply_supported_uniform",
        a_mm=None,
        phi_creep=0.0,
        deflection_limit_profile="aesthetic_open_view",
        available_gap_mm=None,
        w_limit_mm=0.3,
        load_duration="long_term",
    )

    deflection_df = _build_deflection_curve_df(result, service_input)

    assert list(deflection_df.columns) == ["Крок", "M, кН·м", "f, мм"]
    assert len(deflection_df) == len(result.curve_points)
    assert deflection_df.iloc[0]["f, мм"] == pytest.approx(0.0)
    assert deflection_df.iloc[-1]["f, мм"] == pytest.approx(abs(result.curve_points[-1].curvature_1_per_m) * 3750.0)


def test_streamlit_app_updates_serviceability_only_after_serviceability_refresh():
    app = AppTest.from_file("streamlit_app.py")

    app.run(timeout=10)
    _activate_workspace(app, "serviceability")
    baseline_capacity = float(_find_metric(app, "Несуча здатність M_Rd, кН·м").value)
    baseline_deflection = float(_find_metric(app, "Розрахунковий прогин f, мм").value)

    _find_widget_by_label(app.number_input, "Розрахунковий проліт l, мм").set_value(9000.0)
    app.run(timeout=10)

    assert len(app.exception) == 0
    assert float(_find_metric(app, "Несуча здатність M_Rd, кН·м").value) == baseline_capacity
    assert float(_find_metric(app, "Розрахунковий прогин f, мм").value) == baseline_deflection

    _find_widget_by_label(app.button, "Оновити перевірку").click()
    app.run(timeout=10)

    assert float(_find_metric(app, "Несуча здатність M_Rd, кН·м").value) == baseline_capacity
    assert float(_find_metric(app, "Розрахунковий прогин f, мм").value) != baseline_deflection


def test_streamlit_app_does_not_mark_default_serviceability_inputs_as_dirty():
    app = AppTest.from_file("streamlit_app.py")

    app.run(timeout=10)
    serviceability_root = _activate_workspace(app, "serviceability")
    toolbar_html = _find_markdown_containing(app, 'data-role="workspace-nav-toolbar"')

    assert len(app.exception) == 0
    assert "Параметри II ГГС змінено, але ще не застосовано" not in toolbar_html
    assert not any(
        "Є незастосовані зміни для перевірки прогинів і тріщиностійкості." in warning.value
        for warning in app.warning
    )
    assert any(
        "Показано актуальні результати перевірки II групи граничних станів." in success.value
        for success in app.success
    )
    assert _has_markdown_containing_in_node(serviceability_root, 'data-role="serviceability-scheme-showcase"')


def test_streamlit_app_renders_label_offset_controls_for_annotated_charts():
    app = AppTest.from_file("streamlit_app.py")

    app.run(timeout=10)

    slider_labels = [getattr(widget, "label", None) for widget in app.slider]

    assert len(app.exception) == 0
    assert "Зсув X підпису граничної деформації бетону" in slider_labels
    assert "Зсув Y підпису граничної деформації бетону" in slider_labels
    assert "Зсув X підпису εmax бетону" in slider_labels
    assert "Зсув Y підпису εmax бетону" in slider_labels
    assert "Зсув X підпису граничної деформації верхньої арматури" in slider_labels
    assert "Зсув Y підпису εmax верхньої арматури" in slider_labels
    assert "Зсув X підпису граничної деформації нижньої арматури" in slider_labels
    assert "Зсув Y підпису εmax нижньої арматури" in slider_labels
    _activate_workspace(app, "serviceability")
    slider_labels = [getattr(widget, "label", None) for widget in app.slider]
    assert "Зсув X підпису нормативної межі прогину" in slider_labels
    assert "Зсув Y підпису нормативної межі прогину" in slider_labels


def test_streamlit_app_keeps_label_offset_controls_independent():
    app = AppTest.from_file("streamlit_app.py")

    app.run(timeout=10)
    _find_widget_by_label(app.slider, "Зсув X підпису граничної деформації бетону").set_value(18)
    app.run(timeout=10)

    assert len(app.exception) == 0
    assert _find_widget_by_label(app.slider, "Зсув X підпису граничної деформації бетону").value == 18
    assert _find_widget_by_label(app.slider, "Зсув X підпису εmax бетону").value == 0
    assert _find_widget_by_label(app.slider, "Зсув X підпису граничної деформації верхньої арматури").value == 0
    _activate_workspace(app, "serviceability")
    assert _find_widget_by_label(app.slider, "Зсув X підпису нормативної межі прогину").value == 0


def test_streamlit_app_applies_manual_concrete_limit_offset_without_moving_epsilon_max():
    app = AppTest.from_file("streamlit_app.py")

    app.run(timeout=10)
    section_tab = app.main
    baseline_chart = _first_vega_spec_in_node_block_with_subheader(section_tab, "Момент-деформація бетону")
    _find_widget_by_label(app.slider, "Зсув X підпису граничної деформації бетону").set_value(22)
    _find_widget_by_label(app.slider, "Зсув Y підпису граничної деформації бетону").set_value(-11)
    app.run(timeout=10)
    shifted_chart = _first_vega_spec_in_node_block_with_subheader(app.main, "Момент-деформація бетону")

    baseline_limit_text_layer = baseline_chart["layer"][5]["mark"]
    shifted_limit_text_layer = shifted_chart["layer"][5]["mark"]
    baseline_max_text_layer = baseline_chart["layer"][9]["mark"]
    shifted_max_text_layer = shifted_chart["layer"][9]["mark"]

    assert shifted_limit_text_layer["dx"] == baseline_limit_text_layer["dx"] + 22
    assert shifted_limit_text_layer["dy"] == baseline_limit_text_layer["dy"] - 11
    assert shifted_max_text_layer["dx"] == baseline_max_text_layer["dx"]
    assert shifted_max_text_layer["dy"] == baseline_max_text_layer["dy"]


def test_streamlit_app_applies_manual_concrete_epsilon_max_offset_without_moving_limit():
    app = AppTest.from_file("streamlit_app.py")

    app.run(timeout=10)
    section_tab = app.main
    baseline_chart = _first_vega_spec_in_node_block_with_subheader(section_tab, "Момент-деформація бетону")
    _find_widget_by_label(app.slider, "Зсув X підпису εmax бетону").set_value(-19)
    _find_widget_by_label(app.slider, "Зсув Y підпису εmax бетону").set_value(13)
    app.run(timeout=10)
    shifted_chart = _first_vega_spec_in_node_block_with_subheader(app.main, "Момент-деформація бетону")

    baseline_limit_text_layer = baseline_chart["layer"][5]["mark"]
    shifted_limit_text_layer = shifted_chart["layer"][5]["mark"]
    baseline_max_text_layer = baseline_chart["layer"][9]["mark"]
    shifted_max_text_layer = shifted_chart["layer"][9]["mark"]

    assert shifted_limit_text_layer["dx"] == baseline_limit_text_layer["dx"]
    assert shifted_limit_text_layer["dy"] == baseline_limit_text_layer["dy"]
    assert shifted_max_text_layer["dx"] == baseline_max_text_layer["dx"] - 19
    assert shifted_max_text_layer["dy"] == baseline_max_text_layer["dy"] + 13


def test_streamlit_app_defaults_slider_to_last_curve_point():
    catalog = load_material_catalog()
    draft = default_draft_inputs()
    section = build_section_input_from_draft(draft, catalog)
    result = solve_bending_capacity(section, catalog, outer_steps=int(draft["outer_steps"]))

    app = AppTest.from_file("streamlit_app.py")
    app.run(timeout=10)

    slider = _find_widget_by_label(app.slider, "Активна точка кривої")

    assert len(app.exception) == 0
    assert slider.value == result.curve_points[-1].step_index


def test_streamlit_app_hides_missing_first_form_panel_for_default_second_form_state():
    app = AppTest.from_file("streamlit_app.py")

    app.run(timeout=10)

    drawing = _find_markdown_containing(app, "<svg")

    assert len(app.exception) == 0
    assert "2-га форма рівноваги" in drawing
    assert "1-ша форма рівноваги" not in drawing
    assert "форма не реалізована для цього набору даних" not in drawing
    assert drawing.count('data-role="form-panel"') == 1
    hero = _find_markdown_containing(app, 'data-role="hero-banner"')
    assert 'data-role="hero-author-list"' in hero
    assert "Іванейко М.М." in hero
    assert "Іванейко В.М." in hero
    assert not _has_markdown_containing(app, 'data-role="author-section"')


def test_streamlit_app_keeps_last_active_result_until_recalculate():
    app = AppTest.from_file("streamlit_app.py")

    app.run(timeout=10)
    baseline = float(_find_metric(app, "Несуча здатність M_Rd, кН·м").value)
    _find_widget_by_label(app.number_input, "Висота перерізу h, мм").set_value(400.0)
    app.run(timeout=10)

    assert len(app.exception) == 0
    rebar_preview = _find_dataframe(app, ["Шар", "Грань", "a, мм", "z, мм", "n, шт.", "d, мм", "Клас", "Статус"])
    assert rebar_preview.iloc[1]["z, мм"] == 380.0
    assert float(_find_metric(app, "Несуча здатність M_Rd, кН·м").value) == baseline
    assert _find_widget_by_label(app.button, "Перерахувати").disabled is False
    assert any("останнього застосованого" in info.value for info in app.info)
    assert not _has_markdown_containing(app, 'data-role="point-chip-grid"')

    _find_widget_by_label(app.button, "Перерахувати").click()
    app.run(timeout=10)

    assert float(_find_metric(app, "Несуча здатність M_Rd, кН·м").value) != baseline
    assert _has_markdown_containing(app, 'data-role="point-chip-grid"')


def test_streamlit_app_updates_outer_steps_only_after_recalculate():
    app = AppTest.from_file("streamlit_app.py")

    app.run(timeout=10)
    baseline_curve_df = _find_dataframe(
        app,
        ["Крок", "M, кН·м", "κ, 1/м", "ε_c,top, 10^-5", "ε_s,top, 10^-5", "ε_s,bot, 10^-5"],
    )
    baseline_rows = len(baseline_curve_df)

    _find_widget_by_label(app.number_input, "Кількість кроків розрахунку").set_value(8)
    app.run(timeout=10)

    dirty_curve_df = _find_dataframe(
        app,
        ["Крок", "M, кН·м", "κ, 1/м", "ε_c,top, 10^-5", "ε_s,top, 10^-5", "ε_s,bot, 10^-5"],
    )
    assert len(app.exception) == 0
    assert len(dirty_curve_df) == baseline_rows
    assert _find_widget_by_label(app.button, "Перерахувати").disabled is False

    _find_widget_by_label(app.button, "Перерахувати").click()
    app.run(timeout=10)

    updated_curve_df = _find_dataframe(
        app,
        ["Крок", "M, кН·м", "κ, 1/м", "ε_c,top, 10^-5", "ε_s,top, 10^-5", "ε_s,bot, 10^-5"],
    )
    assert len(updated_curve_df) < baseline_rows
    assert len(updated_curve_df) <= 8


def test_streamlit_reference_mode_hides_strengthening_controls_and_uses_single_concrete_row():
    app = AppTest.from_file("streamlit_app.py")

    app.run(timeout=10)
    _find_widget_by_label(app.radio, "Режим бетонного шару").set_value("Без шару")
    app.run(timeout=10)
    _find_widget_by_label(app.selectbox, "Клас бетону еталонної плити").set_value("C25/30")
    app.run(timeout=10)

    concrete_preview = _find_dataframe(app, ["Шар", "b, мм", "h, мм", "Клас", "Статус"])

    assert len(app.exception) == 0
    assert not any(getattr(widget, "label", None) == "Клас бетону верхнього шару" for widget in app.selectbox)
    assert not any(
        getattr(widget, "label", None) == "Товщина верхнього шару бетону h1, мм" for widget in app.number_input
    )
    assert _find_widget_by_label(app.selectbox, "Клас бетону еталонної плити").value == "C25/30"
    assert concrete_preview["Шар"].tolist() == ["B1"]
    assert _find_widget_by_label(app.number_input, "Висота перерізу h, мм").value == 60.0
    assert concrete_preview.iloc[0]["h, мм"] == pytest.approx(60.0)
    assert concrete_preview.iloc[0]["Клас"] == "C25/30"


def test_streamlit_layer_toggle_preserves_hidden_strengthening_parameters():
    app = AppTest.from_file("streamlit_app.py")

    app.run(timeout=10)
    _find_widget_by_label(app.selectbox, "Клас бетону верхнього шару").set_value("C50/60")
    _find_widget_by_label(app.number_input, "Товщина верхнього шару бетону h1, мм").set_value(25.0)
    app.run(timeout=10)
    _find_widget_by_label(app.radio, "Режим бетонного шару").set_value("Без шару")
    app.run(timeout=10)
    assert _find_widget_by_label(app.number_input, "Висота перерізу h, мм").value == 60.0
    _find_widget_by_label(app.radio, "Режим бетонного шару").set_value("З шаром")
    app.run(timeout=10)

    assert len(app.exception) == 0
    assert _find_widget_by_label(app.number_input, "Висота перерізу h, мм").value == 120.0
    assert _find_widget_by_label(app.selectbox, "Клас бетону верхнього шару").value == "C50/60"
    assert _find_widget_by_label(app.number_input, "Товщина верхнього шару бетону h1, мм").value == 25.0


def test_streamlit_reference_mode_uses_single_rebar_input_and_single_rebar_results_panel():
    app = AppTest.from_file("streamlit_app.py")

    app.run(timeout=10)
    _find_widget_by_label(app.radio, "Режим бетонного шару").set_value("Без шару")
    app.run(timeout=10)

    rebar_preview = _find_dataframe(app, ["Шар", "Грань", "a, мм", "z, мм", "n, шт.", "d, мм", "Клас", "Статус"])
    assert not any(getattr(widget, "label", None) == "Грань шару арматури 2" for widget in app.selectbox)
    assert rebar_preview["Шар"].tolist() == ["A1"]
    assert rebar_preview["Грань"].tolist() == ["Нижня"]

    _find_widget_by_label(app.button, "Перерахувати").click()
    app.run(timeout=10)

    drawing = _find_markdown_containing(app, "<svg")
    force_panel = _find_markdown_containing(app, 'data-role="force-card-grid"')
    summary_table = _find_dataframe(app, ["Крок", "M, кН·м", "κ, 1/м", "ε_c,top, 10^-5", "ε_s, 10^-5"])

    assert len(app.exception) == 0
    assert any(getattr(subheader, "value", None) == "Момент-деформація арматури" for subheader in app.subheader)
    assert not any(getattr(subheader, "value", None) == "Момент-деформація верхньої арматури" for subheader in app.subheader)
    assert not any(getattr(subheader, "value", None) == "Момент-деформація нижньої арматури" for subheader in app.subheader)
    assert "A2:" not in drawing
    assert "A2" not in force_panel
    assert list(summary_table.columns) == ["Крок", "M, кН·м", "κ, 1/м", "ε_c,top, 10^-5", "ε_s, 10^-5"]


def test_streamlit_layer_toggle_preserves_hidden_top_rebar_parameters():
    app = AppTest.from_file("streamlit_app.py")

    app.run(timeout=10)
    _find_widget_by_label(app.number_input, "Відстань a1, мм").set_value(35.0)
    _find_widget_by_label(app.number_input, "Кількість n1, шт.").set_value(6)
    app.run(timeout=10)
    _find_widget_by_label(app.radio, "Режим бетонного шару").set_value("Без шару")
    app.run(timeout=10)
    _find_widget_by_label(app.radio, "Режим бетонного шару").set_value("З шаром")
    app.run(timeout=10)

    assert len(app.exception) == 0
    assert _find_widget_by_label(app.number_input, "Відстань a1, мм").value == 35.0
    assert _find_widget_by_label(app.number_input, "Кількість n1, шт.").value == 6


def test_streamlit_app_disables_recalculate_for_invalid_draft():
    app = AppTest.from_file("streamlit_app.py")

    app.run(timeout=10)
    baseline = float(_find_metric(app, "Несуча здатність M_Rd, кН·м").value)
    _find_widget_by_label(app.number_input, "Товщина верхнього шару бетону h1, мм").set_value(600.0)
    app.run(timeout=10)

    assert len(app.exception) == 0
    assert _find_widget_by_label(app.button, "Перерахувати").disabled is True
    assert any("h1" in error.value for error in app.error)
    assert float(_find_metric(app, "Несуча здатність M_Rd, кН·м").value) == baseline


def test_streamlit_app_updates_preview_when_rebar_face_changes():
    app = AppTest.from_file("streamlit_app.py")

    app.run(timeout=10)
    _find_widget_by_label(app.selectbox, "Грань шару арматури 2").set_value("Верхня")
    app.run(timeout=10)

    assert len(app.exception) == 0
    rebar_preview = _find_dataframe(app, ["Шар", "Грань", "a, мм", "z, мм", "n, шт.", "d, мм", "Клас", "Статус"])
    assert rebar_preview.iloc[1]["Грань"] == "Верхня"
    assert rebar_preview.iloc[1]["z, мм"] == 20.0
    drawing = _find_markdown_containing(app, "<svg")
    assert "z2 = 20.0 мм" in drawing


def test_streamlit_app_uses_constant_width_without_individual_layer_controls():
    app = AppTest.from_file("streamlit_app.py")

    app.run(timeout=10)

    labels = [widget.label for widget in app.number_input]
    checkbox_labels = [widget.label for widget in app.checkbox]
    assert len(app.exception) == 0
    assert "Окремі ширини шарів" not in checkbox_labels
    assert "Ширина верхнього шару бетону b1, мм" not in labels
    assert "Ширина нижнього шару бетону b2, мм" not in labels
    drawing = _find_markdown_containing(app, "<svg")
    assert "b = 500.0 мм" in drawing


def test_streamlit_app_updates_section_drawing_when_geometry_changes():
    app = AppTest.from_file("streamlit_app.py")

    app.run(timeout=10)
    initial_drawing = _find_markdown_containing(app, "<svg")
    _find_widget_by_label(app.number_input, "Товщина верхнього шару бетону h1, мм").set_value(50.0)
    app.run(timeout=10)

    updated_drawing = _find_markdown_containing(app, "<svg")

    assert len(app.exception) == 0
    assert "h1 = 60.0 мм" in initial_drawing
    assert "h1 = 50.0 мм" in updated_drawing
    assert "h2 = 70.0 мм" in updated_drawing


def test_streamlit_app_allows_section_height_from_60_mm_and_bottom_layer_below_100_mm():
    app = AppTest.from_file("streamlit_app.py")

    app.run(timeout=10)
    _find_widget_by_label(app.number_input, "Висота перерізу h, мм").set_value(60.0)
    _find_widget_by_label(app.number_input, "Товщина верхнього шару бетону h1, мм").set_value(10.0)
    app.run(timeout=10)

    concrete_preview = _find_dataframe(app, ["Шар", "b, мм", "h, мм", "Клас", "Статус"])
    updated_drawing = _find_markdown_containing(app, "<svg")

    assert len(app.exception) == 0
    assert _find_widget_by_label(app.number_input, "Висота перерізу h, мм").value == 60.0
    assert concrete_preview["h, мм"].tolist() == [pytest.approx(10.0), pytest.approx(50.0)]
    assert "h2 = 50.0 мм" in updated_drawing
    assert not any("h2" in error.value for error in app.error)


def test_streamlit_app_hides_result_overlays_for_dirty_draft_until_recalculate():
    app = AppTest.from_file("streamlit_app.py")

    app.run(timeout=10)
    _find_widget_by_label(app.number_input, "Висота перерізу h, мм").set_value(400.0)
    app.run(timeout=10)

    dirty_drawing = _find_markdown_containing(app, "<svg")

    assert "Активна точка" not in dirty_drawing
    assert "N_A1 =" not in dirty_drawing
    assert dirty_drawing.count('data-role="form-placeholder"') == 2
    assert "оверлеї" in dirty_drawing
    assert not _has_markdown_containing(app, 'data-role="point-chip-grid"')

    _find_widget_by_label(app.button, "Перерахувати").click()
    app.run(timeout=10)

    active_drawing = _find_markdown_containing(app, "<svg")

    assert "Активна точка" in active_drawing
    assert "N_A1 =" in active_drawing
    assert _has_markdown_containing(app, 'data-role="point-chip-grid"')


def test_streamlit_app_updates_current_point_panel_when_curve_point_changes():
    app = AppTest.from_file("streamlit_app.py")

    app.run(timeout=10)
    baseline_panel = _find_markdown_containing(app, 'data-role="point-chip-grid"')

    _find_widget_by_label(app.slider, "Активна точка кривої").set_value(1)
    app.run(timeout=10)

    assert len(app.exception) == 0
    updated_panel = _find_markdown_containing(app, 'data-role="point-chip-grid"')
    assert baseline_panel != updated_panel
    assert 'data-role="point-value-step">1<' in updated_panel


def test_streamlit_app_renders_experimental_empty_state_without_loaded_series():
    app = AppTest.from_file("streamlit_app.py")

    app.run(timeout=10)
    experimental_root = _activate_workspace(app, "experimental")
    indic_toggle = _find_widget_by_label(app.checkbox, "Показати Indic")
    dic_toggle = _find_widget_by_label(app.checkbox, "Показати DIC")

    assert len(app.exception) == 0
    assert _has_markdown_containing_in_node(experimental_root, 'data-role="experimental-empty-state"')
    assert _has_markdown_containing_in_node(
        experimental_root,
        "Завантажте Indic або DIC, щоб побачити накладання на теорію.",
    )
    assert not _has_markdown_containing_in_node(
        experimental_root,
        "Експериментальні дані для Indic, DIC не завантажені. Показано лише доступні криві.",
    )
    assert indic_toggle.disabled is True
    assert dic_toggle.disabled is True
    assert indic_toggle.value is False
    assert dic_toggle.value is False


def test_build_experimental_partial_state_message_mentions_only_missing_slots():
    from streamlit_app import _build_experimental_partial_state_message

    message = _build_experimental_partial_state_message(
        available_slots=["indic", "dic"],
        loaded_slots=["indic"],
    )

    assert message == "Зараз показано лише доступні криві: Indic. Завантажте DIC, щоб додати їх до порівняння."



def test_streamlit_app_renders_indic_and_dic_experimental_charts_and_toggles_visibility():
    app = AppTest.from_file("streamlit_app.py")
    app.session_state["experimental_indic_workbook_bytes"] = _build_experimental_workbook_bytes(
        {
            "M_f": [["f, мм", "M, кН·м"], [0.0, 0.0], [3.5, 12.0], [30.0, 40.0]],
            "M_eps_c": [["ε_c,top, 10^-5", "M, кН·м"], [0.0, 0.0], [120.0, 10.0], [320.0, 30.0]],
            "M_eps_s_top": [["ε_s, 10^-5", "M, кН·м"], [0.0, 0.0], [-120.0, 10.0], [-420.0, 30.0]],
            "M_eps_s_bot": [["ε_s, 10^-5", "M, кН·м"], [0.0, 0.0], [180.0, 10.0], [540.0, 30.0]],
        }
    )
    app.session_state["experimental_indic_workbook_name"] = "indic.xlsx"
    app.session_state["experimental_dic_workbook_bytes"] = _build_experimental_workbook_bytes(
        {
            "M_f": [["f, мм", "M, кН·м"], [0.0, 0.0], [3.0, 11.0], [31.0, 39.5]],
            "M_eps_c": [["ε_c,top, 10^-5", "M, кН·м"], [0.0, 0.0], [130.0, 9.0], [310.0, 28.0]],
            "M_eps_s_top": [["ε_s, 10^-5", "M, кН·м"], [0.0, 0.0], [-130.0, 9.0], [-430.0, 28.0]],
            "M_eps_s_bot": [["ε_s, 10^-5", "M, кН·м"], [0.0, 0.0], [170.0, 9.0], [520.0, 28.0]],
        }
    )
    app.session_state["experimental_dic_workbook_name"] = "dic.xlsx"

    app.run(timeout=10)
    experimental_tab = _activate_workspace(app, "experimental")

    concrete_experimental_tab = _find_tab(app, "M-εc")

    assert len(app.exception) == 0
    assert _has_markdown_containing_in_node(experimental_tab, "indic.xlsx")
    assert _has_markdown_containing_in_node(experimental_tab, "dic.xlsx")
    assert _has_markdown_containing_in_node(experimental_tab, "M_f")
    assert _has_markdown_containing_in_node(experimental_tab, "M_eps_c")
    assert _has_markdown_containing_in_node(experimental_tab, "M_eps_s_top")
    assert _has_markdown_containing_in_node(experimental_tab, "M_eps_s_bot")
    assert _has_markdown_containing_in_node(experimental_tab, 'data-role="experimental-explanation-card"')
    assert _has_markdown_containing_in_node(experimental_tab, 'data-role="experimental-reference-card"')
    assert _has_markdown_containing_in_node(experimental_tab, "M_theory(f_u)")
    assert _has_markdown_containing_in_node(experimental_tab, "M_Indic(f_u)")
    assert _has_markdown_containing_in_node(experimental_tab, "M_DIC(f_u)")
    assert _has_markdown_containing_in_node(experimental_tab, "ε_c(theory @ f_u)")
    assert _has_markdown_containing_in_node(experimental_tab, "ε_s(theory @ f_u)")
    assert _has_markdown_containing_in_node(experimental_tab, "Defl_Indic")
    assert _has_markdown_containing_in_node(experimental_tab, "Defl_DIC")

    concrete_chart_with_theory = _first_vega_spec_in_node(concrete_experimental_tab)
    concrete_layers_with_theory = concrete_chart_with_theory["layer"]
    indic_layers = [layer for layer in concrete_layers_with_theory if layer.get("mark", {}).get("color") == "#b45309"]
    dic_layers = [layer for layer in concrete_layers_with_theory if layer.get("mark", {}).get("color") == "#059669"]
    theory_layers = [
        layer
        for layer in concrete_layers_with_theory
        if layer.get("mark", {}).get("color") == "#1d4ed8"
    ]

    assert len(indic_layers) == 1
    assert len(dic_layers) == 1
    assert len(theory_layers) >= 1
    assert len(concrete_layers_with_theory) >= 7

    _find_widget_by_label(app.checkbox, "Показати DIC").set_value(False)
    app.run(timeout=10)

    concrete_chart_without_dic = _first_vega_spec_in_node(_find_tab(app, "M-εc"))
    concrete_layers_without_dic = concrete_chart_without_dic["layer"]

    assert len(concrete_layers_without_dic) < len(concrete_layers_with_theory)
    assert any(layer.get("mark", {}).get("color") == "#b45309" for layer in concrete_layers_without_dic)
    assert not any(layer.get("mark", {}).get("color") == "#059669" for layer in concrete_layers_without_dic)
    assert _count_markdown_containing(app, "Порівняння для Indic") >= 1
    assert _count_markdown_containing(app, "Порівняння для DIC") >= 1
    assert _has_markdown_containing_in_node(experimental_tab, 'data-role="experimental-fu-summary"')
    comparison_df = _find_dataframe(app, ["x_exp", "M_exp, кН·м", "M_theory_interp, кН·м", "ΔM, кН·м", "ΔM, %", "Статус"])
    assert comparison_df.iloc[0]["Статус"] == "OK"


def test_streamlit_app_places_experimental_fu_summary_below_graph_tabs():
    app = AppTest.from_file("streamlit_app.py")
    app.session_state["experimental_indic_workbook_bytes"] = _build_experimental_workbook_bytes(
        {
            "M_f": [["f, мм", "M, кН·м"], [0.0, 0.0], [3.5, 12.0], [30.0, 40.0]],
            "M_eps_c": [["ε_c,top, 10^-5", "M, кН·м"], [0.0, 0.0], [120.0, 10.0], [320.0, 30.0]],
            "M_eps_s_top": [["ε_s, 10^-5", "M, кН·м"], [0.0, 0.0], [-120.0, 10.0], [-420.0, 30.0]],
            "M_eps_s_bot": [["ε_s, 10^-5", "M, кН·м"], [0.0, 0.0], [180.0, 10.0], [540.0, 30.0]],
        }
    )
    app.session_state["experimental_indic_workbook_name"] = "indic.xlsx"
    app.session_state["experimental_dic_workbook_bytes"] = _build_experimental_workbook_bytes(
        {
            "M_f": [["f, мм", "M, кН·м"], [0.0, 0.0], [3.0, 11.0], [31.0, 39.5]],
            "M_eps_c": [["ε_c,top, 10^-5", "M, кН·м"], [0.0, 0.0], [130.0, 9.0], [310.0, 28.0]],
            "M_eps_s_top": [["ε_s, 10^-5", "M, кН·м"], [0.0, 0.0], [-130.0, 9.0], [-430.0, 28.0]],
            "M_eps_s_bot": [["ε_s, 10^-5", "M, кН·м"], [0.0, 0.0], [170.0, 9.0], [520.0, 28.0]],
        }
    )
    app.session_state["experimental_dic_workbook_name"] = "dic.xlsx"

    app.run(timeout=10)
    experimental_tab = _activate_workspace(app, "experimental")

    chart_index, chart_child_index = _block_and_child_index_with_subheader_in_node(experimental_tab, "Діаграма M-f")
    summary_index, summary_child_index = _block_and_child_index_with_markdown_in_node(
        experimental_tab,
        'data-role="experimental-fu-summary"',
    )
    summary_html = _find_markdown_containing_in_node(experimental_tab, 'data-role="experimental-fu-summary"')

    assert len(app.exception) == 0
    assert summary_index > chart_index or (summary_index == chart_index and summary_child_index > chart_child_index)
    assert "M_theory(f_u)" in summary_html
    assert "M_Indic(f_u)" in summary_html
    assert "M_DIC(f_u)" in summary_html


def test_apply_uploaded_experimental_workbook_keeps_single_sheet_detection_pending_per_slot():
    from streamlit_app import _apply_uploaded_experimental_workbook

    workbook_bytes = _build_positioned_experimental_workbook_bytes(
        {
            "A1": "Діаграма M-f",
            "A2": "f, мм",
            "B2": "M, кН·м",
            "A3": 0.0,
            "B3": 0.0,
            "A4": 3.0,
            "B4": 12.0,
            "D1": "Момент-деформація бетону",
            "D2": "ε_c,top, 10^-5",
            "E2": "M, кН·м",
            "D3": 0.0,
            "E3": 0.0,
            "D4": 130.0,
            "E4": 24.0,
            "A8": "Верхня арматура",
            "A9": "ε_s, 10^-5",
            "B9": "M, кН·м",
            "A10": -240.0,
            "B10": 16.0,
            "A11": -120.0,
            "B11": 7.0,
            "D8": "Нижня арматура",
            "D9": "ε_s, 10^-5",
            "E9": "M, кН·м",
            "D10": 0.0,
            "E10": 0.0,
            "D11": 210.0,
            "E11": 29.0,
        }
    )
    state: dict[str, object] = {}

    _apply_uploaded_experimental_workbook(state, workbook_bytes=workbook_bytes, file_name="lab.xlsx", slot="dic")

    assert state["experimental_dic_pending_workbook_name"] == "lab.xlsx"
    assert state["experimental_dic_pending_workbook_bytes"] == workbook_bytes
    assert "experimental_dic_workbook_bytes" not in state
    inspection = state["experimental_dic_pending_inspection"]
    assert inspection.status == "ready"
    assert inspection.dataset is not None


def test_apply_uploaded_experimental_workbook_auto_accepts_legacy_template_for_indic_and_dic():
    from streamlit_app import _apply_uploaded_experimental_workbook

    workbook_bytes = _build_experimental_workbook_bytes(
        {
            "M_f": [["f, мм", "M, кН·м"], [0.0, 0.0], [3.5, 12.0]],
            "M_eps_c": [["ε_c,top, 10^-5", "M, кН·м"], [0.0, 0.0], [120.0, 10.0]],
        }
    )
    state: dict[str, object] = {}

    _apply_uploaded_experimental_workbook(
        state,
        workbook_bytes=workbook_bytes,
        file_name="indic-template.xlsx",
        slot="indic",
    )
    _apply_uploaded_experimental_workbook(
        state,
        workbook_bytes=workbook_bytes,
        file_name="dic-template.xlsx",
        slot="dic",
    )

    assert state["experimental_indic_workbook_name"] == "indic-template.xlsx"
    assert state["experimental_indic_workbook_bytes"] == workbook_bytes
    assert state["experimental_dic_workbook_name"] == "dic-template.xlsx"
    assert state["experimental_dic_workbook_bytes"] == workbook_bytes
    assert "experimental_indic_pending_workbook_bytes" not in state
    assert "experimental_dic_pending_workbook_bytes" not in state
    assert "experimental_indic_pending_inspection" not in state
    assert "experimental_dic_pending_inspection" not in state


def test_streamlit_app_renders_pending_experimental_preview_and_confirms_it_per_slot():
    app = AppTest.from_file("streamlit_app.py")
    app.session_state["experimental_dic_pending_workbook_bytes"] = _build_positioned_experimental_workbook_bytes(
        {
            "A1": "Діаграма M-f",
            "A2": "f, мм",
            "B2": "M, кН·м",
            "A3": 0.0,
            "B3": 0.0,
            "A4": 3.0,
            "B4": 12.0,
            "D1": "Момент-деформація бетону",
            "D2": "ε_c,top, 10^-5",
            "E2": "M, кН·м",
            "D3": 0.0,
            "E3": 0.0,
            "D4": 130.0,
            "E4": 24.0,
            "A8": "Верхня арматура",
            "A9": "ε_s, 10^-5",
            "B9": "M, кН·м",
            "A10": -240.0,
            "B10": 16.0,
            "A11": -120.0,
            "B11": 7.0,
            "D8": "Нижня арматура",
            "D9": "ε_s, 10^-5",
            "E9": "M, кН·м",
            "D10": 0.0,
            "E10": 0.0,
            "D11": 210.0,
            "E11": 29.0,
        }
    )
    app.session_state["experimental_dic_pending_workbook_name"] = "lab.xlsx"

    app.run(timeout=10)
    experimental_tab = _activate_workspace(app, "experimental")

    pending_concrete_chart = _first_vega_spec_in_node(_find_tab(app, "M-εc"))

    assert len(app.exception) == 0
    assert _has_markdown_containing_in_node(experimental_tab, "Очікує підтвердження")
    assert _has_markdown_containing_in_node(experimental_tab, "lab.xlsx")
    assert _has_markdown_containing_in_node(experimental_tab, "DIC")
    assert _has_markdown_containing_in_node(experimental_tab, "A2:B4")
    assert _has_markdown_containing_in_node(experimental_tab, "M_eps_s_bot")
    assert _has_markdown_containing_in_node(experimental_tab, "Графіки ще не показуються")
    assert _has_markdown_containing_in_node(experimental_tab, "Підтвердити DIC")
    assert not any(layer.get("mark", {}).get("color") == "#b45309" for layer in pending_concrete_chart["layer"])

    _find_widget_by_label(app.button, "Підтвердити DIC").click()
    app.run(timeout=10)
    experimental_tab = _activate_workspace(app, "experimental")

    concrete_chart = _first_vega_spec_in_node(_find_tab(app, "M-εc"))

    assert len(app.exception) == 0
    assert not _has_markdown_containing_in_node(experimental_tab, "Очікує підтвердження")
    assert _has_markdown_containing_in_node(experimental_tab, "Активний файл: lab.xlsx")
    assert any(layer.get("mark", {}).get("color") == "#059669" for layer in concrete_chart["layer"])


def test_streamlit_app_shows_incomplete_pending_experimental_message_and_disables_confirm():
    app = AppTest.from_file("streamlit_app.py")
    app.session_state["experimental_indic_pending_workbook_bytes"] = _build_positioned_experimental_workbook_bytes(
        {
            "A1": "Діаграма M-f",
            "A2": "f, мм",
            "B2": "M, кН·м",
            "A3": 0.0,
            "B3": 0.0,
            "A4": 3.0,
            "B4": 12.0,
            "D1": "Момент-деформація бетону",
            "D2": "ε_c,top, 10^-5",
            "E2": "M, кН·м",
            "D3": 0.0,
            "E3": 0.0,
            "D4": 130.0,
            "E4": 24.0,
            "D8": "Нижня арматура",
            "D9": "ε_s, 10^-5",
            "E9": "M, кН·м",
            "D10": 0.0,
            "E10": 0.0,
            "D11": 210.0,
            "E11": 29.0,
        }
    )
    app.session_state["experimental_indic_pending_workbook_name"] = "missing-top.xlsx"

    app.run(timeout=10)
    experimental_tab = _activate_workspace(app, "experimental")

    confirm_button = _find_widget_by_label(app.button, "Підтвердити Indic")

    assert len(app.exception) == 0
    assert _has_markdown_containing_in_node(experimental_tab, "missing-top.xlsx")
    assert _has_markdown_containing_in_node(experimental_tab, "Відсутні: M_eps_s_top")
    assert _has_markdown_containing_in_node(experimental_tab, "Графіки ще не показуються")
    assert confirm_button.disabled is True


def test_streamlit_reference_mode_uses_only_indic_controls_in_experiment_tab():
    app = AppTest.from_file("streamlit_app.py")
    app.session_state["experimental_indic_workbook_bytes"] = _build_experimental_workbook_bytes(
        {
            "M_f": [["f, мм", "M, кН·м"], [0.0, 0.0], [3.5, 12.0], [30.0, 40.0]],
            "M_eps_c": [["ε_c,top, 10^-5", "M, кН·м"], [0.0, 0.0], [120.0, 10.0], [320.0, 30.0]],
            "M_eps_s_bot": [["ε_s, 10^-5", "M, кН·м"], [0.0, 0.0], [180.0, 10.0], [540.0, 30.0]],
        }
    )
    app.session_state["experimental_indic_workbook_name"] = "reference-indic.xlsx"

    app.run(timeout=10)
    _find_widget_by_label(app.radio, "Режим бетонного шару").set_value("Без шару")
    app.run(timeout=10)
    experimental_tab = _activate_workspace(app, "experimental")

    assert len(app.exception) == 0
    assert _has_markdown_containing_in_node(experimental_tab, "reference-indic.xlsx")
    assert _find_widget_by_label(app.checkbox, "Показати Indic").value is True
    assert _has_markdown_containing_in_node(experimental_tab, "Defl_Indic")
    assert not _has_markdown_containing_in_node(experimental_tab, "Defl_DIC")
    summary_html = _find_markdown_containing_in_node(experimental_tab, 'data-role="experimental-fu-summary"')
    assert "M_Indic(f_u)" in summary_html
    assert "M_DIC(f_u)" not in summary_html


def test_automation_mode_renders_hidden_payload_markers(monkeypatch):
    monkeypatch.setenv("RC_BENDING_AUTOMATION", "1")
    app = AppTest.from_file("streamlit_app.py")

    app.run(timeout=10)

    assert len(app.exception) == 0
    assert _has_markdown_containing(app, 'data-role="automation-section-payload"')
    assert _has_markdown_containing(app, 'data-role="automation-serviceability-payload"')
    assert _has_markdown_containing(app, 'data-role="automation-experimental-payload"')

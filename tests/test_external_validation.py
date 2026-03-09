import os
from pathlib import Path

from openpyxl import load_workbook
import pytest

from rc_bending.materials import load_material_catalog


FIXTURE_DIR = Path(__file__).with_name("fixtures")


def test_default_case_maps_to_expected_eurocodeapplied_submission():
    from rc_bending.external_validation import build_default_validation_case, build_eurocodeapplied_submission

    case = build_default_validation_case(load_material_catalog())
    submission = build_eurocodeapplied_submission(case)

    assert case.name == "canonical_homogeneous_rectangular_section"
    assert case.section.section_height_mm == pytest.approx(500.0)
    assert submission.fields["Calculation.AnalysisType"] == "1"
    assert submission.fields["Calculation.h"] == "0.5"
    assert submission.fields["Calculation.b"] == "0.3"
    assert submission.fields["Calculation.fck"] == "25"
    assert submission.fields["Calculation.fyk"] == "500"
    assert submission.fields["Calculation.acc"] == "0.85"
    assert submission.fields["Calculation.gammaC"] == "1.25"
    assert submission.fields["Calculation.gammaS"] == "1.2"
    assert submission.fields["Calculation.epsilonudRule"] == "6"
    assert submission.fields["Calculation.nBars1"] == "2"
    assert submission.fields["Calculation.Phi1"] == "20"
    assert submission.fields["Calculation.cPrime1"] == "0.04"
    assert submission.fields["Calculation.nBars2"] == "2"
    assert submission.fields["Calculation.Phi2"] == "20"
    assert submission.fields["Calculation.cPrime2"] == "0.04"
    assert submission.fields["Calculation.nBars3"] == "0"


def test_parse_eurocodeapplied_response_normalizes_positive_bending_row():
    from rc_bending.external_validation import parse_eurocodeapplied_response

    html = (FIXTURE_DIR / "eurocodeapplied_mrd_response.html").read_text(encoding="utf-8")

    result = parse_eurocodeapplied_response(html, source_url="https://eurocodeapplied.com/example")

    assert result.source_url == "https://eurocodeapplied.com/example"
    assert result.moment_kNm == pytest.approx(113.16)
    assert result.neutral_axis_depth_mm == pytest.approx(63.0)
    assert result.top_concrete_strain == pytest.approx(0.00159)
    assert result.bottom_concrete_strain == pytest.approx(-0.01101)


def test_compare_validation_results_marks_run_pass_when_metrics_are_within_tolerance():
    from rc_bending.external_validation import (
        ExternalValidationResult,
        LocalValidationResult,
        ValidationCase,
        compare_validation_results,
    )
    from rc_bending.models import ConcreteLayerInput, RebarLayerInput, SectionInput

    case = ValidationCase(
        name="demo",
        description="demo",
        section=SectionInput(
            section_height_mm=500.0,
            concrete_layers=(
                ConcreteLayerInput(width_mm=300.0, height_mm=250.0, concrete_class="C25/30"),
                ConcreteLayerInput(width_mm=300.0, height_mm=250.0, concrete_class="C25/30"),
            ),
            rebar_layers=(
                RebarLayerInput(z_mm=40.0, area_mm2=628.32, steel_class="B500"),
                RebarLayerInput(z_mm=460.0, area_mm2=628.32, steel_class="B500"),
            ),
        ),
        moment_tolerance_ratio=0.01,
        secondary_tolerance_ratio=0.10,
    )
    local = LocalValidationResult(
        moment_kNm=113.57,
        neutral_axis_depth_mm=57.84,
        top_concrete_strain=0.00155,
        bottom_concrete_strain=-0.01185,
        axial_residual_kN=0.01,
    )
    external = ExternalValidationResult(
        source_url="https://eurocodeapplied.com/example",
        moment_kNm=113.16,
        neutral_axis_depth_mm=63.0,
        top_concrete_strain=0.00159,
        bottom_concrete_strain=-0.01101,
        raw_row_index=1,
    )

    comparison = compare_validation_results(case, local, external)

    assert comparison.passed is True
    assert comparison.metrics["moment_kNm"].passed is True
    assert comparison.metrics["neutral_axis_depth_mm"].passed is True
    assert comparison.metrics["top_concrete_strain"].passed is True
    assert comparison.metrics["bottom_concrete_strain"].passed is True


def test_compare_validation_results_marks_run_failed_when_moment_exceeds_tolerance():
    from rc_bending.external_validation import (
        ExternalValidationResult,
        LocalValidationResult,
        ValidationCase,
        compare_validation_results,
    )
    from rc_bending.models import ConcreteLayerInput, RebarLayerInput, SectionInput

    case = ValidationCase(
        name="demo",
        description="demo",
        section=SectionInput(
            section_height_mm=500.0,
            concrete_layers=(
                ConcreteLayerInput(width_mm=300.0, height_mm=250.0, concrete_class="C25/30"),
                ConcreteLayerInput(width_mm=300.0, height_mm=250.0, concrete_class="C25/30"),
            ),
            rebar_layers=(
                RebarLayerInput(z_mm=40.0, area_mm2=628.32, steel_class="B500"),
                RebarLayerInput(z_mm=460.0, area_mm2=628.32, steel_class="B500"),
            ),
        ),
        moment_tolerance_ratio=0.01,
        secondary_tolerance_ratio=0.10,
    )
    local = LocalValidationResult(
        moment_kNm=113.57,
        neutral_axis_depth_mm=57.84,
        top_concrete_strain=0.00155,
        bottom_concrete_strain=-0.01185,
        axial_residual_kN=0.01,
    )
    external = ExternalValidationResult(
        source_url="https://eurocodeapplied.com/example",
        moment_kNm=120.0,
        neutral_axis_depth_mm=63.0,
        top_concrete_strain=0.00159,
        bottom_concrete_strain=-0.01101,
        raw_row_index=1,
    )

    comparison = compare_validation_results(case, local, external)

    assert comparison.passed is False
    assert comparison.metrics["moment_kNm"].passed is False
    assert comparison.metrics["neutral_axis_depth_mm"].passed is True


def test_validation_workbook_contains_required_sheets_and_status_cells(tmp_path):
    from rc_bending.external_validation import (
        EvidenceArtifact,
        ExternalValidationResult,
        LocalValidationResult,
        ValidationCase,
        ValidationRunResult,
        build_validation_workbook_bytes,
        compare_validation_results,
    )
    from rc_bending.models import ConcreteLayerInput, RebarLayerInput, SectionInput

    case = ValidationCase(
        name="demo",
        description="demo",
        section=SectionInput(
            section_height_mm=500.0,
            concrete_layers=(
                ConcreteLayerInput(width_mm=300.0, height_mm=250.0, concrete_class="C25/30"),
                ConcreteLayerInput(width_mm=300.0, height_mm=250.0, concrete_class="C25/30"),
            ),
            rebar_layers=(
                RebarLayerInput(z_mm=40.0, area_mm2=628.32, steel_class="B500"),
                RebarLayerInput(z_mm=460.0, area_mm2=628.32, steel_class="B500"),
            ),
        ),
        moment_tolerance_ratio=0.01,
        secondary_tolerance_ratio=0.10,
    )
    local = LocalValidationResult(
        moment_kNm=113.57,
        neutral_axis_depth_mm=57.84,
        top_concrete_strain=0.00155,
        bottom_concrete_strain=-0.01185,
        axial_residual_kN=0.01,
    )
    external = ExternalValidationResult(
        source_url="https://eurocodeapplied.com/example",
        moment_kNm=113.16,
        neutral_axis_depth_mm=63.0,
        top_concrete_strain=0.00159,
        bottom_concrete_strain=-0.01101,
        raw_row_index=1,
    )
    comparison = compare_validation_results(case, local, external)
    run_result = ValidationRunResult(
        case=case,
        local_result=local,
        external_result=external,
        comparison=comparison,
        evidence=(
            EvidenceArtifact(
                label="Screenshot",
                path=r"C:\теорія розрахунок\output\playwright\eurocodeapplied.png",
                description="Result page screenshot.",
            ),
        ),
        status="passed",
        message="Validation completed.",
        executed_at="2026-03-09T10:15:00+02:00",
    )

    workbook_bytes = build_validation_workbook_bytes(run_result)
    output_path = tmp_path / "external_validation_report.xlsx"
    output_path.write_bytes(workbook_bytes)
    workbook = load_workbook(output_path)

    assert workbook.sheetnames == [
        "Summary",
        "CaseInputs",
        "LocalResults",
        "EurocodeAppliedResults",
        "Comparison",
        "Evidence",
    ]
    assert workbook["Summary"]["B2"].value == "passed"
    assert workbook["Summary"]["B3"].value == "Validation completed."
    assert workbook["Evidence"]["B2"].value.endswith("eurocodeapplied.png")


def test_select_mcp_profile_processes_filters_only_target_profile():
    from rc_bending.external_validation import select_mcp_profile_processes

    rows = [
        {"ProcessId": 101, "Name": "chrome.exe", "CommandLine": r"chrome.exe --user-data-dir=C:\Users\Muhailo\AppData\Local\ms-playwright\mcp-chrome"},
        {"ProcessId": 102, "Name": "chrome.exe", "CommandLine": r"chrome.exe --user-data-dir=C:\Users\Muhailo\AppData\Local\Google\Chrome\User Data"},
        {"ProcessId": 103, "Name": "msedge.exe", "CommandLine": r"msedge.exe --user-data-dir=C:\Users\Muhailo\AppData\Local\ms-playwright\mcp-chrome"},
        {"ProcessId": 104, "Name": "chrome.exe", "CommandLine": None},
    ]

    matches = select_mcp_profile_processes(rows, r"C:\Users\Muhailo\AppData\Local\ms-playwright\mcp-chrome")

    assert [row["ProcessId"] for row in matches] == [101]


@pytest.mark.skipif(os.getenv("LIVE_EUROCODEAPPLIED_VALIDATION") != "1", reason="Set LIVE_EUROCODEAPPLIED_VALIDATION=1 to run live smoke validation.")
def test_live_eurocodeapplied_validation_smoke(tmp_path):
    from rc_bending.external_validation import build_validation_workbook_bytes, run_default_validation

    base_dir = Path(__file__).resolve().parents[1]
    run_result, _ = run_default_validation(
        base_dir,
        load_material_catalog(),
        evidence_paths=(),
        execute_preflight=False,
    )
    output_path = tmp_path / "live_external_validation_report.xlsx"
    output_path.write_bytes(build_validation_workbook_bytes(run_result))

    assert run_result.status == "passed"
    assert output_path.exists()

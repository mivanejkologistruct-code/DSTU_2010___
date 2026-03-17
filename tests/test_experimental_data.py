from io import BytesIO
from pathlib import Path

from openpyxl import Workbook, load_workbook
import pandas as pd
import pytest

from rc_bending.experimental_data import (
    ExperimentalSeries,
    build_comparison_table,
    build_experimental_template_workbook_bytes,
    inspect_experimental_workbook,
    load_experimental_dataset,
)


def _build_workbook_bytes(sheet_rows: dict[str, list[list[object]]]) -> bytes:
    workbook = Workbook()
    first_sheet = workbook.active
    first_title = next(iter(sheet_rows))
    first_sheet.title = first_title

    for sheet_index, (sheet_name, rows) in enumerate(sheet_rows.items()):
        sheet = first_sheet if sheet_index == 0 else workbook.create_sheet(sheet_name)
        for row in rows:
            sheet.append(row)

    buffer = BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def _build_positioned_workbook_bytes(
    sheet_cells: dict[str, dict[str, object]],
    *,
    sheet_name: str = "Experiment",
) -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = sheet_name

    for cell_ref, value in sheet_cells.items():
        sheet[cell_ref] = value

    buffer = BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def test_build_experimental_template_workbook_contains_supported_sheets_and_headers():
    workbook = load_workbook(BytesIO(build_experimental_template_workbook_bytes()))

    assert workbook.sheetnames == ["M_f", "M_eps_c", "M_eps_s_top", "M_eps_s_bot"]
    assert [workbook["M_f"]["A1"].value, workbook["M_f"]["B1"].value] == ["f, мм", "M, кН·м"]
    assert [workbook["M_eps_c"]["A1"].value, workbook["M_eps_c"]["B1"].value] == ["ε_c,top, 10^-5", "M, кН·м"]
    assert [workbook["M_eps_s_top"]["A1"].value, workbook["M_eps_s_top"]["B1"].value] == ["ε_s, 10^-5", "M, кН·м"]
    assert [workbook["M_eps_s_bot"]["A1"].value, workbook["M_eps_s_bot"]["B1"].value] == ["ε_s, 10^-5", "M, кН·м"]


def test_load_experimental_dataset_reads_all_supported_sheets():
    workbook_bytes = _build_workbook_bytes(
        {
            "M_f": [["f, мм", "M, кН·м"], [0.0, 0.0], [4.2, 38.5]],
            "M_eps_c": [["ε_c,top, 10^-5", "M, кН·м"], [0.0, 0.0], [110.0, 42.0]],
            "M_eps_s_top": [["ε_s, 10^-5", "M, кН·м"], [-300.0, 18.0], [-120.0, 6.0]],
            "M_eps_s_bot": [["ε_s, 10^-5", "M, кН·м"], [0.0, 0.0], [180.0, 24.0]],
        }
    )

    dataset = load_experimental_dataset(workbook_bytes)

    assert set(dataset.series) == {"M_f", "M_eps_c", "M_eps_s_top", "M_eps_s_bot"}
    assert dataset.series["M_f"].x_field == "f, мм"
    assert list(dataset.series["M_f"].data.columns) == ["f, мм", "M, кН·м"]
    assert dataset.series["M_eps_s_top"].data.iloc[0]["ε_s, 10^-5"] == pytest.approx(-300.0)


def test_load_experimental_dataset_allows_partial_supported_set_and_ignores_unknown_sheets():
    workbook_bytes = _build_workbook_bytes(
        {
            "M_f": [["f, мм", "M, кН·м"], [0.0, 0.0], [2.5, 28.0]],
            "Notes": [["comment"], ["lab run 3"]],
        }
    )

    dataset = load_experimental_dataset(workbook_bytes)

    assert set(dataset.series) == {"M_f"}
    assert dataset.ignored_sheets == ("Notes",)


def test_load_experimental_dataset_rejects_workbook_without_supported_sheets():
    workbook_bytes = _build_workbook_bytes({"Notes": [["comment"], ["no charts here"]]})

    with pytest.raises(ValueError, match="жодного підтримуваного"):
        load_experimental_dataset(workbook_bytes)


def test_load_experimental_dataset_rejects_sheet_with_missing_required_columns():
    workbook_bytes = _build_workbook_bytes({"M_f": [["f, мм", "Moment"], [0.0, 0.0], [1.0, 12.0]]})

    with pytest.raises(ValueError, match="M_f"):
        load_experimental_dataset(workbook_bytes)


def test_load_experimental_dataset_rejects_non_numeric_rows_but_ignores_fully_blank_rows():
    workbook_bytes = _build_workbook_bytes(
        {
            "M_eps_c": [
                ["ε_c,top, 10^-5", "M, кН·м"],
                [None, None],
                [0.0, 0.0],
                ["bad", 15.0],
            ]
        }
    )

    with pytest.raises(ValueError, match="M_eps_c"):
        load_experimental_dataset(workbook_bytes)


def test_load_experimental_dataset_requires_at_least_two_valid_points_per_sheet():
    workbook_bytes = _build_workbook_bytes({"M_eps_s_bot": [["ε_s, 10^-5", "M, кН·м"], [150.0, 20.0]]})

    with pytest.raises(ValueError, match="щонайменше 2"):
        load_experimental_dataset(workbook_bytes)


def test_load_experimental_dataset_drops_fully_blank_rows_inside_sheet():
    workbook_bytes = _build_workbook_bytes(
        {
            "M_f": [
                ["f, мм", "M, кН·м"],
                [None, None],
                [0.0, 0.0],
                [3.0, 22.0],
                [None, None],
            ]
        }
    )

    dataset = load_experimental_dataset(workbook_bytes)

    assert len(dataset.series["M_f"].data) == 2


def test_build_comparison_table_interpolates_theoretical_moment_for_experimental_points():
    theory_df = pd.DataFrame(
        {
            "Крок": [1, 2, 3],
            "f, мм": [0.0, 5.0, 10.0],
            "M, кН·м": [0.0, 50.0, 100.0],
        }
    )
    experimental_series = ExperimentalSeries(
        sheet_name="M_f",
        x_field="f, мм",
        data=pd.DataFrame({"f, мм": [2.5, 5.0, 12.0], "M, кН·м": [24.0, 55.0, 120.0]}),
    )

    comparison_df = build_comparison_table(theory_df, experimental_series)

    assert list(comparison_df.columns) == [
        "x_exp",
        "M_exp, кН·м",
        "M_theory_interp, кН·м",
        "ΔM, кН·м",
        "ΔM, %",
        "Статус",
    ]
    assert comparison_df.iloc[0]["M_theory_interp, кН·м"] == pytest.approx(25.0)
    assert comparison_df.iloc[0]["ΔM, кН·м"] == pytest.approx(-1.0)
    assert comparison_df.iloc[1]["M_theory_interp, кН·м"] == pytest.approx(50.0)
    assert comparison_df.iloc[1]["ΔM, %"] == pytest.approx(10.0)
    assert pd.isna(comparison_df.iloc[2]["M_theory_interp, кН·м"])
    assert comparison_df.iloc[2]["Статус"] == "Поза діапазоном теорії"


def test_build_comparison_table_handles_descending_theory_x_axis_for_top_rebar():
    theory_df = pd.DataFrame(
        {
            "Крок": [1, 2, 3],
            "ε_s, 10^-5": [-120.0, -220.0, -320.0],
            "M, кН·м": [10.0, 25.0, 40.0],
        }
    )
    experimental_series = ExperimentalSeries(
        sheet_name="M_eps_s_top",
        x_field="ε_s, 10^-5",
        data=pd.DataFrame({"ε_s, 10^-5": [-170.0, -320.0], "M, кН·м": [16.0, 39.0]}),
    )

    comparison_df = build_comparison_table(theory_df, experimental_series)

    assert comparison_df.iloc[0]["M_theory_interp, кН·м"] == pytest.approx(17.5)
    assert comparison_df.iloc[1]["M_theory_interp, кН·м"] == pytest.approx(40.0)


def test_build_comparison_table_leaves_percentage_empty_when_theoretical_moment_is_zero():
    theory_df = pd.DataFrame(
        {
            "Крок": [1, 2],
            "ε_c,top, 10^-5": [0.0, 100.0],
            "M, кН·м": [0.0, 50.0],
        }
    )
    experimental_series = ExperimentalSeries(
        sheet_name="M_eps_c",
        x_field="ε_c,top, 10^-5",
        data=pd.DataFrame({"ε_c,top, 10^-5": [0.0], "M, кН·м": [2.0]}),
    )

    comparison_df = build_comparison_table(theory_df, experimental_series)

    assert comparison_df.iloc[0]["M_theory_interp, кН·м"] == pytest.approx(0.0)
    assert pd.isna(comparison_df.iloc[0]["ΔM, %"])


def test_inspect_experimental_workbook_detects_four_labeled_tables_on_single_sheet():
    workbook_bytes = _build_positioned_workbook_bytes(
        {
            "A1": "Діаграма M-f",
            "A2": "f, мм",
            "B2": "M, кН·м",
            "A3": 0.0,
            "B3": 0.0,
            "A4": 3.5,
            "B4": 14.0,
            "D1": "Момент-деформація бетону",
            "D2": "ε_c,top, 10^-5",
            "E2": "M, кН·м",
            "D3": 0.0,
            "E3": 0.0,
            "D4": 150.0,
            "E4": 26.0,
            "A8": "Верхня арматура",
            "A9": "ε_s, 10^-5",
            "B9": "M, кН·м",
            "A10": -250.0,
            "B10": 18.0,
            "A11": -120.0,
            "B11": 8.0,
            "D8": "Нижня арматура",
            "D9": "ε_s, 10^-5",
            "E9": "M, кН·м",
            "D10": 0.0,
            "E10": 0.0,
            "D11": 220.0,
            "E11": 30.0,
        }
    )

    inspection = inspect_experimental_workbook(workbook_bytes)

    assert inspection.status == "ready"
    assert inspection.dataset is not None
    assert inspection.missing_graphs == ()
    assert inspection.conflicts == {}
    assert {block.target_graph for block in inspection.detected_blocks} == {
        "M_f",
        "M_eps_c",
        "M_eps_s_top",
        "M_eps_s_bot",
    }
    assert inspection.dataset.series["M_f"].data.iloc[1]["f, мм"] == pytest.approx(3.5)
    assert inspection.dataset.series["M_eps_s_top"].data.iloc[0]["ε_s, 10^-5"] == pytest.approx(-250.0)


def test_inspect_experimental_workbook_reports_missing_graph_for_incomplete_single_sheet():
    workbook_bytes = _build_positioned_workbook_bytes(
        {
            "A1": "Діаграма M-f",
            "A2": "f, мм",
            "B2": "M, кН·м",
            "A3": 0.0,
            "B3": 0.0,
            "A4": 2.0,
            "B4": 10.0,
            "D1": "Момент-деформація бетону",
            "D2": "ε_c,top, 10^-5",
            "E2": "M, кН·м",
            "D3": 0.0,
            "E3": 0.0,
            "D4": 120.0,
            "E4": 21.0,
            "A8": "Нижня арматура",
            "A9": "ε_s, 10^-5",
            "B9": "M, кН·м",
            "A10": 0.0,
            "B10": 0.0,
            "A11": 180.0,
            "B11": 25.0,
        }
    )

    inspection = inspect_experimental_workbook(workbook_bytes)

    assert inspection.status == "incomplete"
    assert inspection.dataset is None
    assert inspection.missing_graphs == ("M_eps_s_top",)
    assert inspection.conflicts == {}


def test_inspect_experimental_workbook_reports_conflict_when_two_blocks_match_same_graph():
    workbook_bytes = _build_positioned_workbook_bytes(
        {
            "A1": "Діаграма M-f",
            "A2": "f, мм",
            "B2": "M, кН·м",
            "A3": 0.0,
            "B3": 0.0,
            "A4": 2.0,
            "B4": 10.0,
            "G1": "M-f повтор",
            "G2": "f, мм",
            "H2": "M, кН·м",
            "G3": 0.0,
            "H3": 0.0,
            "G4": 3.0,
            "H4": 11.0,
            "D1": "Момент-деформація бетону",
            "D2": "ε_c,top, 10^-5",
            "E2": "M, кН·м",
            "D3": 0.0,
            "E3": 0.0,
            "D4": 140.0,
            "E4": 19.0,
            "A8": "Верхня арматура",
            "A9": "ε_s, 10^-5",
            "B9": "M, кН·м",
            "A10": -250.0,
            "B10": 18.0,
            "A11": -120.0,
            "B11": 8.0,
            "D8": "Нижня арматура",
            "D9": "ε_s, 10^-5",
            "E9": "M, кН·м",
            "D10": 0.0,
            "E10": 0.0,
            "D11": 220.0,
            "E11": 30.0,
        }
    )

    inspection = inspect_experimental_workbook(workbook_bytes)

    assert inspection.status == "conflict"
    assert inspection.dataset is None
    assert "M_f" in inspection.conflicts


def test_load_experimental_dataset_rejects_single_sheet_block_with_non_numeric_row():
    workbook_bytes = _build_positioned_workbook_bytes(
        {
            "A1": "Діаграма M-f",
            "A2": "f, мм",
            "B2": "M, кН·м",
            "A3": 0.0,
            "B3": 0.0,
            "A4": 3.5,
            "B4": 14.0,
            "D1": "Момент-деформація бетону",
            "D2": "ε_c,top, 10^-5",
            "E2": "M, кН·м",
            "D3": 0.0,
            "E3": 0.0,
            "D4": "bad",
            "E4": 26.0,
            "A8": "Верхня арматура",
            "A9": "ε_s, 10^-5",
            "B9": "M, кН·м",
            "A10": -250.0,
            "B10": 18.0,
            "A11": -120.0,
            "B11": 8.0,
            "D8": "Нижня арматура",
            "D9": "ε_s, 10^-5",
            "E9": "M, кН·м",
            "D10": 0.0,
            "E10": 0.0,
            "D11": 220.0,
            "E11": 30.0,
        }
    )

    with pytest.raises(ValueError, match="нечислові або неповні дані"):
        load_experimental_dataset(workbook_bytes)


@pytest.mark.parametrize(
    ("fixture_name", "expected_graph", "expected_cell_range"),
    [
        ("lab_one_sheet_primary.xlsx", "M_f", "A2:B4"),
        ("lab_one_sheet_shifted.xlsx", "M_eps_c", "B3:C5"),
    ],
)
def test_inspect_experimental_workbook_parses_anonymized_fixture_workbooks(
    fixture_name: str,
    expected_graph: str,
    expected_cell_range: str,
):
    fixture_path = Path("tests/fixtures/experimental_workbooks") / fixture_name
    workbook_bytes = fixture_path.read_bytes()

    inspection = inspect_experimental_workbook(workbook_bytes)

    assert inspection.status == "ready"
    assert inspection.dataset is not None
    assert set(inspection.dataset.series) == {"M_f", "M_eps_c", "M_eps_s_top", "M_eps_s_bot"}
    assert any(
        block.target_graph == expected_graph and block.cell_range == expected_cell_range
        for block in inspection.detected_blocks
    )
